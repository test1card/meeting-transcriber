from __future__ import annotations

import json
import tempfile
from pathlib import Path

from docx import Document

from meeting_transcriber.documents import build_all
from meeting_transcriber.jobs import DOCUMENTS


def write_json(path: Path, value: object) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="meeting-documents-smoke-") as temp:
        root = Path(temp)
        job = {
            "input": {"filename": "запись совещания.mp4"},
            "text_processing": {
                "mode": "deepseek",
                "provider": "deepseek",
                "model": "deepseek-v4-flash",
            },
        }
        write_json(root / "job.json", job)
        transcript = {
            "metadata": {"complete": True},
            "transcript": {
                "language": "ru",
                "segments": [
                    {
                        "start": 1161.94,
                        "end": 1168.14,
                        "speaker": "Speaker 5",
                        "text": "Сейчас надо определиться с оборудованием.",
                    },
                    {
                        "start": 1169.44,
                        "end": 1174.64,
                        "speaker": "Speaker 5",
                        "text": "У меня полного понимания нет.",
                    },
                ],
            },
        }
        write_json(root / "stages" / "consensus" / "transcript.json", transcript)
        write_json(
            root / "stages" / "cleaned" / "transcript-cleaned.json",
            transcript,
        )
        write_json(
            root / "stages" / "condensed" / "transcript-condensed.json",
            transcript,
        )
        minutes = root / "stages" / "minutes"
        minutes.mkdir(parents=True)
        (minutes / "minutes-deepseek-v4-flash.md").write_text(
            "# Итоги\n\n"
            "- Определиться с оборудованием.\n"
            "- Уточнить расположение оборудования.\n\n"
            "| Поручение | Ответственный | Срок |\n"
            "|---|---|---|\n"
            "| Уточнить мощность | не указан | не указан |\n",
            encoding="utf-8",
        )

        manifest = build_all(root)
        assert len(manifest) == 4
        assert [item["filename"] for item in manifest] == [
            item["filename"] for item in DOCUMENTS
        ]
        for item in manifest:
            path = Path(item["path"])
            assert path.is_file()
            document = Document(path)
            visible = "\n".join(p.text for p in document.paragraphs)
            assert visible.strip()
        minutes_doc = Document(root / "documents" / DOCUMENTS[0]["filename"])
        assert len(minutes_doc.tables) == 1
        assert minutes_doc.tables[0].cell(0, 0).text == "Поручение"
        assert minutes_doc.tables[0].cell(1, 0).text == "Уточнить мощность"
        full = Document(root / "documents" / DOCUMENTS[1]["filename"])
        visible = "\n".join(p.text for p in full.paragraphs)
        assert "[00:19:21–00:19:28]" in visible
        assert "Speaker 5" in visible
        print(json.dumps({"documents": manifest}, ensure_ascii=False))


if __name__ == "__main__":
    main()
