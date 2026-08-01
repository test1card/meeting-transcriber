from __future__ import annotations

import argparse
import json
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


INK = RGBColor(31, 45, 61)
BLUE = RGBColor(46, 116, 181)
MUTED = RGBColor(96, 108, 122)
PALE = "E8EEF5"
GREEN = "E8F3EC"
WARN = "FFF4D6"


def set_font(run, *, size: float = 11, bold: bool = False, color=INK) -> None:
    run.font.name = "Calibri"
    fonts = run._element.get_or_add_rPr().rFonts
    fonts.set(qn("w:ascii"), "Calibri")
    fonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade(cell, fill: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    element = props.find(qn("w:shd"))
    if element is None:
        element = OxmlElement("w:shd")
        props.append(element)
    element.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    props = cell._tc.get_or_add_tcPr()
    margins = props.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        props.append(margins)
    for edge, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        element = margins.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            margins.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int], indent: int = 120) -> None:
    table.autofit = False
    props = table._tbl.tblPr
    width = props.first_child_found_in("w:tblW")
    if width is None:
        width = OxmlElement("w:tblW")
        props.append(width)
    width.set(qn("w:w"), str(sum(widths)))
    width.set(qn("w:type"), "dxa")
    tbl_ind = props.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        props.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for value in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(value))
        grid.append(column)
    for row in table.rows:
        for cell, value in zip(row.cells, widths):
            tcw = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            tcw.set(qn("w:w"), str(value))
            tcw.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def configure(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(1)
    section.left_margin = section.right_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)

    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, before, after in (
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 13, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = document.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLUE
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(
        header.add_run("Meeting Documents · оценка протокола"),
        size=9,
        color=MUTED,
    )
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(
        footer.add_run("Внутренний рабочий отчёт · 30.07.2026"),
        size=8.5,
        color=MUTED,
    )


def add_title(document: Document) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(12)
    paragraph.paragraph_format.space_after = Pt(4)
    set_font(
        paragraph.add_run("Оценка автоматического протокола совещания"),
        size=23,
        bold=True,
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_after = Pt(14)
    set_font(
        paragraph.add_run(
            "DeepSeek V4 Flash и локальные Gemma: качество, приватность, "
            "время и hardware offload"
        ),
        size=13,
        color=MUTED,
    )


def add_callout(document: Document, title: str, text: str, fill: str = GREEN) -> None:
    table = document.add_table(rows=1, cols=1)
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    shade(cell, fill)
    paragraph = cell.paragraphs[0]
    set_font(paragraph.add_run(title + "\n"), size=11, bold=True)
    set_font(paragraph.add_run(text), size=10.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def add_bullet(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(4)
    paragraph.add_run(text)


def fmt(value: float) -> str:
    return f"{value:.3f}".replace(".", ",")


def gib(value: int | None) -> str:
    return "—" if not value else f"{value / 1024**3:.1f} ГиБ".replace(".", ",")


def find_run(
    runs: list[dict], model: str, profile: str = "p3_human_fidelity"
) -> dict | None:
    return next(
        (
            item
            for item in runs
            if item["model"] == model and item["profile"] == profile
        ),
        None,
    )


def add_main_matrix(document: Document, runs: list[dict]) -> None:
    document.add_heading("Сводная оценка", level=1)
    document.add_paragraph(
        "Human — сходство с рабочим протоколом человека. Grounded — соответствие "
        "тому, что действительно можно подтвердить стенограммой. Показатели "
        "разделены намеренно: человеческий документ содержит ФИО, сроки и параметры "
        "из внешних материалов, которых нет в аудио."
    )
    selected = [
        find_run(runs, model)
        for model in (
            "deepseek-v4-flash",
            "gemma4:12b",
            "gemma4:e4b",
            "gemma4:26b",
            "gemma4:31b",
        )
    ]
    table = document.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = [
        "Модель",
        "Human",
        "Grounded",
        "Обсуждение F1",
        "Решения F1",
        "Числа",
        "Время",
    ]
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(label), size=9, bold=True)
    for item in (value for value in selected if value):
        values = [
            item["model"],
            fmt(item["human_similarity_score"]),
            fmt(item["transcript_grounded_score"]),
            fmt(item["discussed_vs_human"]["f1"]),
            fmt(item["resolved_vs_human"]["f1"]),
            fmt(item["numeric_fidelity_on_matched_items"]),
            f'{item["elapsed_seconds"]:.1f} с'.replace(".", ","),
        ]
        for cell, value in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(value), size=8.5)
    set_table_geometry(table, [2100, 1000, 1100, 1400, 1300, 1000, 1460])


def add_sensitivity(document: Document, runs: list[dict]) -> None:
    document.add_heading("Чувствительность к формулировке prompt", level=1)
    document.add_paragraph(
        "Δ показан относительно минимального P0. Положительное значение означает "
        "улучшение; отрицательное — ухудшение."
    )
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Модель", "Профиль", "Δ Human", "Δ Grounded", "Что менялось"]
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(label), size=9, bold=True)
    for model in ("deepseek-v4-flash", "gemma4:12b", "gemma4:e4b"):
        baseline = find_run(runs, model, "p0_minimal")
        if not baseline:
            continue
        variants = (
            ("p1_secretary_structure", "Одна форма"),
            ("p2_strict_grounding", "Одни запреты"),
            ("p3_human_fidelity", "Форма + категории + gate"),
        )
        for profile, interpretation in variants:
            item = find_run(runs, model, profile)
            if not item:
                continue
            values = [
                model,
                profile[:2].upper(),
                f'{item["human_similarity_score"] - baseline["human_similarity_score"]:+.3f}',
                f'{item["transcript_grounded_score"] - baseline["transcript_grounded_score"]:+.3f}',
                interpretation,
            ]
            for cell, value in zip(table.add_row().cells, values):
                set_font(cell.paragraphs[0].add_run(value.replace(".", ",")), size=8.5)
    set_table_geometry(table, [2100, 900, 1100, 1200, 4060])


def add_privacy_tradeoff(document: Document, runs: list[dict]) -> None:
    document.add_heading("Что теряем при запрете облака", level=1)
    deepseek = find_run(runs, "deepseek-v4-flash")
    local = find_run(runs, "gemma4:12b")
    if deepseek and local:
        delta_h = local["human_similarity_score"] - deepseek["human_similarity_score"]
        delta_g = local["transcript_grounded_score"] - deepseek[
            "transcript_grounded_score"
        ]
        slowdown = local["elapsed_seconds"] / deepseek["elapsed_seconds"]
        add_callout(
            document,
            "Практический privacy trade-off",
            f"Gemma 12B уступила DeepSeek по Human на {abs(delta_h):.3f}, "
            f"по Grounded на {abs(delta_g):.3f} и была медленнее примерно "
            f"в {slowdown:.1f} раза. Разрыв по опоре на стенограмму невелик: "
            "12B остаётся рабочим локальным вариантом, но требует обязательной "
            "проверки дат, ФИО и чисел.",
            WARN,
        )
    add_bullet(
        document,
        "DeepSeek: основной режим для несекретных встреч и лучшего готового протокола.",
    )
    add_bullet(
        document,
        "Gemma 12B: базовый полностью локальный режим для конфиденциальных встреч.",
    )
    add_bullet(
        document,
        "Gemma e4b: не использовать для финального протокола; только быстрый черновик.",
    )


def add_offload(document: Document, runs: list[dict]) -> None:
    heavy = [
        item
        for item in (
            find_run(runs, "gemma4:26b"),
            find_run(runs, "gemma4:31b"),
        )
        if item
    ]
    if not heavy:
        return
    document.add_heading("Тяжёлые локальные модели и offload", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Модель", "Resident", "VRAM", "CPU/RAM offload", "GPU-доля"]
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(label), size=9, bold=True)
    for item in heavy:
        values = [
            item["model"],
            gib(item.get("ollama_model_resident_bytes")),
            gib(item.get("ollama_vram_bytes")),
            gib(item.get("ollama_cpu_offload_bytes")),
            f'{100 * item.get("ollama_gpu_residency_fraction", 0):.1f}%'.replace(
                ".", ","
            ),
        ]
        for cell, value in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(value), size=9)
    set_table_geometry(table, [2200, 1700, 1600, 2200, 1660])


def add_cost(document: Document, runs: list[dict]) -> None:
    item = find_run(runs, "deepseek-v4-flash")
    if not item:
        return
    hit = item.get("prompt_cache_hit_tokens", 0)
    miss = item.get("prompt_cache_miss_tokens", item.get("prompt_tokens", 0))
    output = item.get("completion_tokens", 0)
    cost = hit / 1e6 * 0.0028 + miss / 1e6 * 0.14 + output / 1e6 * 0.28
    document.add_heading("Время и стоимость API", level=1)
    document.add_paragraph(
        f"Лучший DeepSeek P3 занял {item['elapsed_seconds']:.1f} секунды и "
        f"использовал {item.get('prompt_tokens', 0):,} входных и {output:,} "
        f"выходных токенов. По официальным тарифам на 30.07.2026 расчётная "
        f"стоимость прогона — около ${cost:.4f}. Электроэнергия локальных "
        "моделей не учитывается: фиксируются только hardware requirements и время."
    )
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    set_font(
        paragraph.add_run(
            "Источник тарифа: https://api-docs.deepseek.com/quick_start/pricing"
        ),
        size=9,
        color=MUTED,
    )


def add_limitations(document: Document) -> None:
    document.add_heading("Ограничения оценки", level=1)
    add_bullet(
        document,
        "Человеческий протокол — целевой стиль и рабочий ориентир, но не абсолютная истина.",
    )
    add_bullet(
        document,
        "Полные ФИО и многие дедлайны не звучат в записи и требуют карточки встречи.",
    )
    add_bullet(
        document,
        "Часть численных расхождений отражает конфликт аудио и внешних материалов, "
        "а не обязательную галлюцинацию модели.",
    )
    add_bullet(
        document,
        "Автоматические scores диагностические; критические решения и числа "
        "проверяются по evidence timestamps.",
    )


def add_prompt_appendix(document: Document, matrix: dict) -> None:
    profile = matrix["profiles"]["p3_human_fidelity"]
    document.add_page_break()
    document.add_heading("Приложение: выигравший профиль P3", level=1)
    document.add_heading("Extraction prompt", level=2)
    document.add_paragraph(profile["extract"])
    document.add_heading("Synthesis prompt", level=2)
    document.add_paragraph(profile["synthesize"])
    document.add_paragraph(
        "Каждый фактически отправленный system/user prompt и ответ модели сохранён "
        "в runs/<model>/<profile>/prompts. Это позволяет воспроизвести сравнение "
        "и проверить авторство."
    )


def build(matrix_path: Path, output_path: Path) -> None:
    matrix = json.loads(matrix_path.read_text(encoding="utf-8"))
    runs = matrix["runs"]
    document = Document()
    configure(document)
    add_title(document)
    add_callout(
        document,
        "Рекомендация",
        "DeepSeek V4 Flash + P3 — основной режим. Gemma 12B + P3 — рабочий "
        "локальный fallback для секретных данных. e4b не годится для финального "
        "протокола. 26B и 31B оцениваются по фактическому offload и времени.",
    )
    add_main_matrix(document, runs)
    add_sensitivity(document, runs)
    add_privacy_tradeoff(document, runs)
    add_offload(document, runs)
    add_cost(document, runs)
    add_limitations(document)
    add_prompt_appendix(document, matrix)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("matrix", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.matrix.resolve(), args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
