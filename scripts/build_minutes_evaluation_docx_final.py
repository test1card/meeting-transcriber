from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from minutes_docx_automation_section import add_docx_automation
from minutes_vibe_section import add_vibe_matrix

from build_minutes_evaluation_docx import (
    BLUE,
    GREEN,
    INK,
    MUTED,
    PALE,
    WARN,
    add_callout,
    configure,
    set_font,
    set_table_geometry,
    shade,
)


RECOMMENDED_PROFILES = {
    "deepseek-v4-flash": "p3_human_fidelity",
    "gemma4:12b": "p3_human_fidelity",
    "gemma4:26b": "p5_gemma26_tuned",
    "gemma4:31b": "p5_gemma31_tuned",
    "gemma4:e4b": "p3_human_fidelity",
}

MODEL_LABELS = {
    "deepseek-v4-flash": "DeepSeek V4 Flash",
    "gemma4:12b": "Gemma 4 12B",
    "gemma4:26b": "Gemma 4 26B",
    "gemma4:31b": "Gemma 4 31B",
    "gemma4:e4b": "Gemma 4 e4b",
}


def get_run(runs: list[dict[str, Any]], model: str, profile: str) -> dict[str, Any] | None:
    return next(
        (row for row in runs if row["model"] == model and row["profile"] == profile),
        None,
    )


def score(value: float) -> str:
    return f"{value:.3f}".replace(".", ",")


def seconds(value: float) -> str:
    if value < 180:
        return f"{value:.1f} с".replace(".", ",")
    minutes = int(value // 60)
    remainder = int(round(value - minutes * 60))
    return f"{minutes} мин {remainder:02d} с"


def gib(value: int | None) -> str:
    if not value:
        return "—"
    return f"{value / 1024**3:.2f} ГиБ".replace(".", ",")


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    set_font(
        p.add_run("Сравнение моделей для протокола совещания"),
        size=23,
        bold=True,
        color=INK,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(12)
    set_font(
        p.add_run(
            "DeepSeek V4 Flash и локальные Gemma 4: качество, чувствительность "
            "к промпту, время и CPU/GPU offload"
        ),
        size=13,
        color=MUTED,
    )
    for label, value in (
        ("Дата измерений", "30.07.2026"),
        ("Тестовая система", "Ryzen 9 9900X · 64 GB DDR5 · RTX 5070 Ti 16 GB"),
        ("Материал", "русскоязычное рабочее совещание с отдельными английскими терминами"),
        ("Эталон", "человеческий протокол + проверяемая стенограмма"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"{label}: "), size=10.5, bold=True)
        set_font(p.add_run(value), size=10.5)


def add_decision(doc: Document, runs: list[dict[str, Any]]) -> None:
    deepseek = get_run(runs, "deepseek-v4-flash", "p3_human_fidelity")
    local = get_run(runs, "gemma4:12b", "p3_human_fidelity")
    if not deepseek or not local:
        return
    human_rel = 100 * (deepseek["human_similarity_score"] - local["human_similarity_score"]) / deepseek["human_similarity_score"]
    grounded_rel = 100 * (deepseek["transcript_grounded_score"] - local["transcript_grounded_score"]) / deepseek["transcript_grounded_score"]
    numeric_rel = 100 * (deepseek["numeric_fidelity_on_matched_items"] - local["numeric_fidelity_on_matched_items"]) / max(
        deepseek["numeric_fidelity_on_matched_items"], 1e-9
    )
    slowdown = local["elapsed_seconds"] / deepseek["elapsed_seconds"]
    add_callout(
        doc,
        "Решение",
        "Основной режим: DeepSeek V4 Flash с профилем P3 Human Fidelity. "
        "Конфиденциальный локальный режим: Gemma 4 12B с тем же P3 и обязательной "
        "проверкой дат, ФИО и чисел. На этом совещании 12B ниже DeepSeek на "
        f"{human_rel:.1f}% по сходству с человеческим протоколом, на {grounded_rel:.1f}% "
        f"по опоре на стенограмму и на {numeric_rel:.1f}% по точности чисел; она "
        f"медленнее в {slowdown:.1f} раза. 26B и 31B не следует ставить по умолчанию "
        "на 16 GB VRAM до тех пор, пока выигрыш не подтверждён повторными встречами.",
        GREEN,
    )


def add_metric_explanation(doc: Document) -> None:
    doc.add_heading("Как читать результат", level=1)
    doc.add_paragraph(
        "Human similarity измеряет близость к документу, подготовленному человеком. "
        "Transcript grounded измеряет соответствие тому, что можно подтвердить самой "
        "стенограммой. Эти показатели нельзя объединять без оговорок: человек использовал "
        "ФИО, сроки и сведения из рабочих материалов, которых в анонимной аудиозаписи нет."
    )
    doc.add_paragraph(
        "Discuss F1 и Resolve F1 оценивают покрытие обсуждений и поручений относительно "
        "человеческого документа. Numbers проверяет совпадение чисел в сопоставленных "
        "пунктах. Все оценки диагностические: итоговое решение дополнено ручной проверкой "
        "ложных сроков, Speaker-ярлыков, языка и смысловой полноты."
    )


def add_recommended_table(doc: Document, runs: list[dict[str, Any]]) -> None:
    doc.add_heading("Рекомендуемый профиль каждой модели", level=1)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Модель", "Профиль", "Human", "Grounded", "Числа", "Время", "Роль"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=8.5, bold=True)
    roles = {
        "deepseek-v4-flash": "основной",
        "gemma4:12b": "локальный fallback",
        "gemma4:26b": "эксперимент",
        "gemma4:31b": "не рекомендован",
        "gemma4:e4b": "черновик",
    }
    for model in (
        "deepseek-v4-flash",
        "gemma4:12b",
        "gemma4:26b",
        "gemma4:31b",
        "gemma4:e4b",
    ):
        row = get_run(runs, model, RECOMMENDED_PROFILES[model])
        if not row:
            continue
        values = [
            MODEL_LABELS[model],
            row["profile"],
            score(row["human_similarity_score"]),
            score(row["transcript_grounded_score"]),
            score(row["numeric_fidelity_on_matched_items"]),
            seconds(row["elapsed_seconds"]),
            roles[model],
        ]
        for cell, text in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(text), size=8)
    set_table_geometry(table, [1780, 1600, 850, 1000, 800, 1430, 1900])


def add_full_matrix(doc: Document, runs: list[dict[str, Any]]) -> None:
    doc.add_heading("Полная матрица промптов", level=1)
    doc.add_paragraph(
        "P0–P3 — одинаковые контрольные профили для честного сравнения. P5/P6 — "
        "отдельные адаптации под конкретную модель. Улучшение не предполагается заранее: "
        "неудачный tuned-профиль остаётся в таблице как результат чувствительности."
    )
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["Модель", "Профиль", "Human", "Grounded", "Discuss F1", "Resolve F1", "Числа", "Время"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=7.8, bold=True)
    for row in sorted(runs, key=lambda item: (item["model"], item["profile"])):
        values = [
            MODEL_LABELS.get(row["model"], row["model"]),
            row["profile"],
            score(row["human_similarity_score"]),
            score(row["transcript_grounded_score"]),
            score(row["discussed_vs_human"]["f1"]),
            score(row["resolved_vs_human"]["f1"]),
            score(row["numeric_fidelity_on_matched_items"]),
            seconds(row["elapsed_seconds"]),
        ]
        for cell, text in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(text), size=7.4)
    set_table_geometry(table, [1550, 1650, 820, 900, 1050, 1050, 750, 1590])


def add_local_analysis(doc: Document, runs: list[dict[str, Any]]) -> None:
    doc.add_heading("Что реально дают 12B, 26B и 31B", level=1)
    doc.add_heading("Gemma 4 12B", level=2)
    doc.add_paragraph(
        "P3 оказался лучшим сбалансированным вариантом. Модель сохранила структуру и "
        "значительную часть содержания, но выдумала срок 17.08.2024 и поэтому требует "
        "детерминированного удаления Speaker-ярлыков и ручной проверки владельцев, дат и чисел. "
        "Секрет практического использования: не просить 12B «быть строже» общими словами — "
        "короткий tuned-профиль P5 ухудшил Human и Grounded. Работает подробный P3 плюс "
        "последующая валидация полей."
    )
    doc.add_heading("Gemma 4 26B", level=2)
    doc.add_paragraph(
        "P2 получил высокий автоматический Grounded, но достиг этого за счёт абстрактных "
        "формулировок без чисел и включил Speaker 1–6 в участников. Tuned P5 убрал Speaker-ярлыки "
        "и вернул параметры, но смешал русский и английский и добавил сомнительного владельца. "
        "То есть 26B не показала стабильного преимущества над 12B, достаточного для установки "
        "по умолчанию на 16 GB VRAM."
    )
    doc.add_heading("Gemma 4 31B", level=2)
    doc.add_paragraph(
        "Из-за крупного CPU-offload модель очень медленна. Даже если отдельный профиль даёт "
        "высокий score, решение принимается только после проверки чисел, Speaker-ярлыков, "
        "сроков и фактической полноты. Время полного цикла — самостоятельный критерий пригодности "
        "для офисного рабочего места."
    )


def add_offload(doc: Document, runs: list[dict[str, Any]]) -> None:
    doc.add_heading("Измеренный offload на тестовой системе", level=1)
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Модель", "Параметры", "Квантизация", "Resident", "VRAM", "CPU/RAM"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=8.5, bold=True)
    details = {
        "gemma4:26b": ("25,8B", "Q4_K_M"),
        "gemma4:31b": ("31,3B", "Q4_K_M"),
    }
    for model in ("gemma4:26b", "gemma4:31b"):
        row = next(
            (
                item
                for item in runs
                if item["model"] == model and item.get("ollama_model_resident_bytes")
            ),
            None,
        )
        if not row:
            continue
        values = [
            MODEL_LABELS[model],
            details[model][0],
            details[model][1],
            gib(row.get("ollama_model_resident_bytes")),
            gib(row.get("ollama_vram_bytes")),
            gib(row.get("ollama_cpu_offload_bytes")),
        ]
        for cell, text in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(text), size=8.3)
    set_table_geometry(table, [1800, 1100, 1300, 1700, 1700, 1760])
    doc.add_paragraph(
        "Локальный API-тариф равен USD 0. Стоимость электроэнергии намеренно не "
        "рассчитывается; учитываются только требования к hardware и измеренное время."
    )


def add_api_cost(doc: Document, runs: list[dict[str, Any]]) -> None:
    row = get_run(runs, "deepseek-v4-flash", "p3_human_fidelity")
    if not row:
        return
    hit = row.get("prompt_cache_hit_tokens", 0)
    miss = row.get("prompt_cache_miss_tokens", row.get("prompt_tokens", 0))
    output = row.get("completion_tokens", 0)
    cost = hit / 1e6 * 0.0028 + miss / 1e6 * 0.14 + output / 1e6 * 0.28
    doc.add_heading("DeepSeek: время и фактическая оценка стоимости", level=1)
    doc.add_paragraph(
        f"P3 занял {seconds(row['elapsed_seconds'])}; вход — {row.get('prompt_tokens', 0):,} "
        f"токенов, выход — {output:,} токенов. Расчётная стоимость этого прогона по тарифам "
        f"DeepSeek на дату измерения — около ${cost:.4f}."
    )
    p = doc.add_paragraph()
    set_font(
        p.add_run("Источник тарифа: https://api-docs.deepseek.com/quick_start/pricing"),
        size=9,
        color=MUTED,
    )


def add_recommendations(doc: Document) -> None:
    doc.add_heading("Точные рекомендации для продукта", level=1)
    recommendations = [
        (
            "1. Несекретные встречи.",
            "DeepSeek V4 Flash + P3. Аудио и распознавание остаются локальными; наружу "
            "отправляется только текст стенограммы после явного согласия.",
        ),
        (
            "2. Конфиденциальные встречи.",
            "Gemma 4 12B + P3. После генерации автоматически удалять Speaker-ярлыки из "
            "участников/владельцев, а даты и числа помечать для проверки секретарём.",
        ),
        (
            "3. 26B на 16 GB VRAM.",
            "Не предлагать как рекомендуемую установку. Оставить в Advanced как "
            "экспериментальный вариант для ручного сравнения.",
        ),
        (
            "4. 31B на 16 GB VRAM.",
            "Не устанавливать по умолчанию: крупный CPU-offload делает обработку слишком "
            "долгой для офисного потока, а стабильного выигрыша качества не подтверждено.",
        ),
        (
            "5. e4b.",
            "Не использовать для финального протокола. Допустим только быстрый черновик, "
            "который не показывается как готовый документ.",
        ),
        (
            "6. Метаданные карточки встречи.",
            "ФИО, дата, место и утверждённая повестка должны вводиться секретарём отдельно. "
            "Это детерминированный слой и самый надёжный способ приблизиться к человеческому документу.",
        ),
    ]
    for label, body in recommendations:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        set_font(p.add_run(label + " "), size=11, bold=True, color=BLUE)
        set_font(p.add_run(body), size=11)


def add_limitations(doc: Document) -> None:
    doc.add_heading("Ограничения и корректная интерпретация", level=1)
    for text in (
        "Одна встреча даёт case study, а не универсальный рейтинг. Решение следует "
        "перепроверить на нескольких типах совещаний.",
        "Человеческий протокол полезен как целевой стиль, но не идеален: он содержит "
        "сведения из внешних материалов и последующей организационной работы.",
        "Высокий автоматический Grounded может быть получен чрезмерно общими фразами. "
        "Поэтому scores всегда сопровождаются ручной проверкой содержательности.",
        "Согласие нескольких STT/LLM не доказывает истинность. Для речи нужны отдельные "
        "human-reference WER/CER, cpWER и DER/JER.",
    ):
        doc.add_paragraph(text)


def add_prompt_audit(doc: Document, matrix: dict[str, Any]) -> None:
    doc.add_page_break()
    doc.add_heading("Приложение: воспроизводимость промптов", level=1)
    doc.add_paragraph(
        "Для каждого запуска сохранены точные system/user messages, JSON schema, полный "
        "ответ модели, время и токены. Архив содержит текст совещания и должен храниться "
        "как конфиденциальный материал; API-ключей в нём нет."
    )
    profile = matrix["profiles"]["p3_human_fidelity"]
    doc.add_heading("Победивший общий профиль P3 — extraction", level=2)
    doc.add_paragraph(profile["extract"])
    doc.add_heading("Победивший общий профиль P3 — synthesis", level=2)
    doc.add_paragraph(profile["synthesize"])


def build(matrix_path: Path, output_path: Path) -> None:
    matrix = json.loads(matrix_path.read_text(encoding="utf-8"))
    runs = matrix["runs"]
    doc = Document()
    configure(doc)
    add_title(doc)
    add_decision(doc, runs)
    add_metric_explanation(doc)
    add_recommended_table(doc, runs)
    add_full_matrix(doc, runs)
    add_local_analysis(doc, runs)
    add_vibe_matrix(doc)
    add_offload(doc, runs)
    add_api_cost(doc, runs)
    add_recommendations(doc)
    add_docx_automation(doc)
    add_limitations(doc)
    add_prompt_audit(doc, matrix)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("matrix", type=Path)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.matrix.resolve(), args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
