from __future__ import annotations

import json
import math
import re
from datetime import datetime
from decimal import Decimal
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from meeting_transcriber.formatting import timestamp


ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "outputs" / "russian-meeting-benchmark"
OUT = BASE / "documents"
OUT.mkdir(parents=True, exist_ok=True)

SKILL_ROOT = (
    Path.home()
    / ".codex"
    / "plugins"
    / "cache"
    / "openai-primary-runtime"
    / "documents"
    / "26.727.11326"
    / "skills"
    / "documents"
)

BLACK = "000000"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
NAVY = "0B2545"
MUTED = "5D6773"
LIGHT_GRAY = "F2F4F7"
BLUE_GRAY = "E8EEF5"
CALLOUT = "F4F6F9"
WHITE = "FFFFFF"
RED = "9B1C1C"
GOLD = "7A5A00"
GREEN = "1F5F42"

PAGE_WIDTH_DXA = 12240
PAGE_HEIGHT_DXA = 15840
MARGIN_DXA = 1440
CONTENT_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
CELL_MARGINS = {"top": 80, "bottom": 80, "start": 120, "end": 120}


def set_run_font(
    run,
    name: str = "Calibri",
    size: float | None = None,
    color: str | None = None,
    bold: bool | None = None,
    italic: bool | None = None,
) -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.rFonts
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.insert(0, fonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attr}"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, margins: dict[str, int] = CELL_MARGINS) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.find(qn("w:tcMar"))
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side in ("top", "start", "bottom", "end"):
        node = tc_mar.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(margins[side]))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")
    cell.width = Inches(width_dxa / 1440)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_table_borders(table, color: str = "C9CED6", size: int = 4) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_table_geometry(table, widths_dxa: list[int], indent_dxa: int = TABLE_INDENT_DXA) -> None:
    if sum(widths_dxa) != CONTENT_WIDTH_DXA:
        raise ValueError(f"Table widths must total {CONTENT_WIDTH_DXA}: {widths_dxa}")
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_WIDTH_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_table_borders(table)


def set_keep_with_next(paragraph, value: bool = True) -> None:
    paragraph.paragraph_format.keep_with_next = value


def add_page_field(paragraph) -> None:
    run = paragraph.add_run("Стр. ")
    set_run_font(run, size=9, color=MUTED)
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    txt = OxmlElement("w:t")
    txt.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, txt, fld_end])


def configure_document(doc: Document, preset: str, running_label: str) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    if preset == "compact_reference_guide":
        pf.line_spacing = 1.25
        heading_values = {
            "Heading 1": (16, BLUE, 18, 10),
            "Heading 2": (13, BLUE, 14, 7),
            "Heading 3": (12, DARK_BLUE, 10, 5),
        }
    else:
        pf.line_spacing = 1.10
        heading_values = {
            "Heading 1": (16, BLUE, 16, 8),
            "Heading 2": (13, BLUE, 12, 6),
            "Heading 3": (12, DARK_BLUE, 8, 4),
        }
    for name, (size, color, before, after) in heading_values.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Named override for dense transcript entries.
    if "Transcript Entry" not in [s.name for s in styles]:
        transcript_style = styles.add_style("Transcript Entry", 1)
    else:
        transcript_style = styles["Transcript Entry"]
    transcript_style.font.name = "Calibri"
    transcript_style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    transcript_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    transcript_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    transcript_style.font.size = Pt(9.5)
    transcript_style.paragraph_format.space_before = Pt(0)
    transcript_style.paragraph_format.space_after = Pt(3)
    transcript_style.paragraph_format.line_spacing = 1.08

    header_p = section.header.paragraphs[0]
    header_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_p.paragraph_format.space_after = Pt(0)
    hrun = header_p.add_run(running_label)
    set_run_font(hrun, size=8.5, color=MUTED, bold=True)

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_p.paragraph_format.space_before = Pt(0)
    footer_p.paragraph_format.space_after = Pt(0)
    add_page_field(footer_p)

    props = doc.core_properties
    props.author = "Meeting Transcription Evaluation"
    props.last_modified_by = "Meeting Transcription Evaluation"
    props.created = datetime(2026, 7, 29, 12, 0, 0)
    props.modified = datetime(2026, 7, 29, 12, 0, 0)


def add_masthead(
    doc: Document,
    kicker: str,
    title: str,
    subtitle: str,
    metadata: Iterable[tuple[str, str]] = (),
) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(kicker.upper())
    set_run_font(r, size=9, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(title)
    set_run_font(r, size=23, color=BLACK, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(14)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(subtitle)
    set_run_font(r, size=12.5, color=MUTED)

    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.0
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=10, color=BLACK, bold=True)
        r = p.add_run(value)
        set_run_font(r, size=10, color=BLACK)
    if list(metadata):
        spacer = doc.add_paragraph()
        spacer.paragraph_format.space_after = Pt(3)


def add_callout(doc: Document, label: str, text: str, fill: str = CALLOUT, color: str = NAVY) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.12)
    ppr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    ppr.append(shd)
    r = p.add_run(f"{label}: ")
    set_run_font(r, size=10.5, bold=True, color=color)
    r = p.add_run(text)
    set_run_font(r, size=10.5, color=BLACK)


def add_inline_runs(paragraph, text: str, size: float = 11) -> None:
    # Support Markdown bold and inline code; strip normal Markdown residue.
    parts = re.split(r"(\*\*.*?\*\*|`.*?`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=size - 0.5, color=DARK_BLUE)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=size)


def add_bullet(doc: Document, text: str, compact: bool = False) -> None:
    p = doc.add_paragraph()
    p.style = doc.styles["Normal"]
    p.paragraph_format.left_indent = Inches(0.5 if not compact else 0.375)
    p.paragraph_format.first_line_indent = Inches(-0.25 if not compact else -0.188)
    p.paragraph_format.space_after = Pt(8 if not compact else 4)
    p.paragraph_format.line_spacing = 1.167 if not compact else 1.25
    marker = p.add_run("•\t")
    set_run_font(marker, size=11)
    add_inline_runs(p, text, 11)


def add_numbered(doc: Document, number: int, text: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.5)
    p.paragraph_format.first_line_indent = Inches(-0.25)
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(f"{number}.\t")
    set_run_font(r, size=11, bold=True, color=BLUE)
    add_inline_runs(p, text, 11)


def add_table(
    doc: Document,
    headers: list[str],
    rows: list[list[str]],
    widths_dxa: list[int],
    compact_font: float = 9.3,
    header_fill: str = LIGHT_GRAY,
    center_columns: set[int] | None = None,
) -> None:
    center_columns = center_columns or set()
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0]
    for idx, value in enumerate(headers):
        p = hdr.cells[idx].paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(value)
        set_run_font(r, size=compact_font, bold=True, color=NAVY)
        set_cell_shading(hdr.cells[idx], header_fill)
    set_repeat_table_header(hdr)
    for row_values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row_values):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = (
                WD_ALIGN_PARAGRAPH.CENTER if idx in center_columns else WD_ALIGN_PARAGRAPH.LEFT
            )
            add_inline_runs(p, str(value), compact_font)
    set_table_geometry(table, widths_dxa)
    after = doc.add_paragraph()
    after.paragraph_format.space_before = Pt(4)
    after.paragraph_format.space_after = Pt(0)


def parse_md_table(lines: list[str], start: int) -> tuple[list[str], list[list[str]], int]:
    block = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        block.append(lines[i].strip())
        i += 1
    parsed = [[cell.strip() for cell in row.strip("|").split("|")] for row in block]
    if len(parsed) < 2:
        return [], [], i
    headers = parsed[0]
    rows = parsed[2:]
    return headers, rows, i


def render_markdown_body(
    doc: Document,
    lines: list[str],
    compact: bool = False,
    skip_first_heading: bool = False,
) -> None:
    i = 0
    skipped = not skip_first_heading
    while i < len(lines):
        raw = lines[i]
        text = raw.strip()
        if not text or text == "---":
            i += 1
            continue
        if text.startswith("#"):
            level = len(text) - len(text.lstrip("#"))
            title = text[level:].strip()
            if not skipped:
                skipped = True
                i += 1
                continue
            style = "Heading 1" if level <= 2 else "Heading 2" if level == 3 else "Heading 3"
            p = doc.add_paragraph(title, style=style)
            set_keep_with_next(p)
            i += 1
            continue
        if text.startswith("|"):
            headers, rows, i = parse_md_table(lines, i)
            widths = [4200, 1900, 1500, 1760] if len(headers) == 4 else [CONTENT_WIDTH_DXA // len(headers)] * len(headers)
            widths[-1] += CONTENT_WIDTH_DXA - sum(widths)
            add_table(
                doc,
                headers,
                rows,
                widths,
                compact_font=8.8,
                header_fill=LIGHT_GRAY if not compact else BLUE_GRAY,
                center_columns={1, 2, 3} if len(headers) == 4 else set(),
            )
            continue
        if re.match(r"^[*+-]\s+", text):
            value = re.sub(r"^[*+-]\s+", "", text)
            add_bullet(doc, value, compact=compact)
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline_runs(p, text, 11)
        i += 1


def build_minutes() -> Path:
    source = BASE / "summary-comparison" / "minutes-deepseek-v4-flash.md"
    lines = source.read_text(encoding="utf-8").splitlines()
    title_line = next(line for line in lines if line.startswith("## "))
    title = title_line[3:].strip()
    start = lines.index(title_line) + 1

    doc = Document()
    configure_document(doc, "standard_business_brief", "ПРОТОКОЛ СОВЕЩАНИЯ")
    add_masthead(
        doc,
        "Рабочий протокол",
        title,
        "Структурированная версия по автоматической стенограмме",
        [
            ("Запись", "74 мин 41 с; основной язык — русский"),
            ("Подготовлено", "29 июля 2026 г."),
            ("Основа", "Консенсус MOSS + Qwen3-ASR + WhisperX; редактура DeepSeek V4 Flash"),
        ],
    )
    add_callout(
        doc,
        "Статус",
        "Рабочий проект для сверки участниками. Имена, цифры, решения и ответственных следует подтвердить по временным меткам в стенограмме.",
        fill="FFF8E8",
        color=GOLD,
    )
    render_markdown_body(doc, lines[start:], compact=False)
    out = OUT / "01_Протокол_совещания.docx"
    doc.save(out)
    return out


TRANSCRIPT_RE = re.compile(
    r"^\[(?P<time>\d{2}:\d{2}:\d{2})\]\s+(?P<speaker>Speaker\s+\d+):\s*(?P<text>.*)$"
)


def build_transcript() -> Path:
    source = BASE / "consensus-deepseek-flash-v2" / "transcript.md"
    lines = source.read_text(encoding="utf-8").splitlines()
    entries = []
    for line in lines:
        match = TRANSCRIPT_RE.match(line.strip())
        if match:
            entries.append(match.groupdict())

    doc = Document()
    configure_document(doc, "compact_reference_guide", "ПОЛНАЯ СТЕНОГРАММА • КОНСЕНСУС STT")
    add_masthead(
        doc,
        "Стенограмма",
        "Полная стенограмма совещания",
        "Русская речь с единичными англоязычными терминами; временные метки относятся к исходной записи",
        [
            ("Продолжительность", "74 мин 41 с"),
            ("Реплики", f"{len(entries)}"),
            ("Голосовые кластеры", "6 (автоматические идентификаторы Speaker 1–6)"),
            ("Дата формирования", "29 июля 2026 г."),
        ],
    )
    add_callout(
        doc,
        "Важно",
        "Speaker 1–6 — автоматические голосовые кластеры, а не подтвержденные личности. Кластеризация выполнена MOSS + ECAPA и требует ручной проверки, особенно для коротких реплик.",
        fill="FFF3F3",
        color=RED,
    )

    current_bucket = -1
    for entry in entries:
        h, m, s = (int(part) for part in entry["time"].split(":"))
        seconds = h * 3600 + m * 60 + s
        bucket = seconds // 600
        if bucket != current_bucket:
            start_m = bucket * 10
            end_m = start_m + 10
            p = doc.add_paragraph(
                f"{start_m // 60:02d}:{start_m % 60:02d}–{end_m // 60:02d}:{end_m % 60:02d}",
                style="Heading 2",
            )
            p.paragraph_format.page_break_before = bucket > 0
            current_bucket = bucket
        p = doc.add_paragraph(style="Transcript Entry")
        p.paragraph_format.keep_together = True
        r = p.add_run(f"[{entry['time']}] ")
        set_run_font(r, size=9.2, color=MUTED)
        r = p.add_run(f"{entry['speaker']}: ")
        set_run_font(r, size=9.5, color=DARK_BLUE, bold=True)
        r = p.add_run(entry["text"])
        set_run_font(r, size=9.5, color=BLACK)

    out = OUT / "02_Полная_стенограмма_с_говорящими.docx"
    doc.save(out)
    return out


def build_cleaned_transcript() -> Path:
    source = BASE / "cleaned-transcript-deepseek-v2" / "transcript-cleaned.json"
    payload = json.loads(source.read_text(encoding="utf-8"), parse_float=Decimal)
    metadata = payload["metadata"]
    segments = payload["transcript"]["segments"]
    visible_segments = [segment for segment in segments if segment["text"]]

    doc = Document()
    configure_document(
        doc,
        "compact_reference_guide",
        "РЕДАКЦИОННАЯ СТЕНОГРАММА • DEEPSEEK CLEANUP",
    )
    add_masthead(
        doc,
        "Редакционная стенограмма",
        "Редакционная стенограмма совещания",
        "Отдельная нечистовая производная для чтения; временные интервалы и метки Speaker N сохранены из cleaned JSON",
        [
            ("Модель очистки", metadata["model"]),
            ("Реплики", f"{metadata['segments']} входных; {metadata['empty_segments']} пустых удалено; {len(visible_segments)} показано"),
            ("Слова", f"{metadata['original_words']} → {metadata['cleaned_words']}"),
            ("Дата формирования", "29 июля 2026 г."),
        ],
    )
    add_callout(
        doc,
        "Внимание: это не дословная стенограмма",
        "DeepSeek V4 Flash удалил слова-паразиты, повторы, ложные старты и ASR-мусор. "
        "Модель не слушала аудио. Факты, числа, имена, формулировки решений и принадлежность "
        "реплик необходимо проверять по документу «02_Полная_стенограмма_с_говорящими.docx» "
        "и исходной аудиозаписи.",
        fill="FFF3F3",
        color=RED,
    )
    add_table(
        doc,
        ["Показатель очистки", "Значение"],
        [
            ["Провайдер / модель", f"{metadata['provider']} / {metadata['model']}"],
            ["Измененные сегменты", f"{metadata['changed_segments']} из {metadata['segments']}"],
            ["Удаленные пустые реплики", f"{metadata['empty_segments']} из {metadata['segments']}"],
            ["Видимые реплики", str(len(visible_segments))],
            ["Слова до / после", f"{metadata['original_words']} / {metadata['cleaned_words']}"],
            ["Сокращение", f"{metadata['word_reduction_fraction'] * 100:.1f}%"],
            ["Время обработки", f"{metadata['elapsed_seconds']} с"],
            ["Стоимость API", f"${metadata['api_cost_usd']}"],
            ["Версия процедуры", metadata["prompt_version"]],
        ],
        [3200, 6160],
        compact_font=9.0,
        header_fill=LIGHT_GRAY,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        f"Происхождение: редакционная очистка канонической консенсусной стенограммы. Получено {metadata['segments']} входных сегментов; {metadata['empty_segments']} реплик без текста удалено; показано {len(visible_segments)} реплик. "
        "Текст ниже воспроизведен из transcript-cleaned.json без дополнительного "
        "перефразирования или суммаризации."
    )
    set_run_font(r, size=9.2, color=MUTED, italic=True)

    current_bucket = -1
    for segment in visible_segments:
        start = segment["start"]
        bucket = int(start) // 600
        if bucket != current_bucket:
            start_m = bucket * 10
            end_m = start_m + 10
            p = doc.add_paragraph(
                f"{start_m // 60:02d}:{start_m % 60:02d}–{end_m // 60:02d}:{end_m % 60:02d}",
                style="Heading 2",
            )
            p.paragraph_format.page_break_before = bucket > 0
            current_bucket = bucket

        p = doc.add_paragraph(style="Transcript Entry")
        p.paragraph_format.keep_together = True
        r = p.add_run(f"[{timestamp(float(start))}] ")
        set_run_font(r, size=9.2, color=MUTED)
        r = p.add_run(f"{segment['speaker']}: ")
        set_run_font(r, size=9.5, color=DARK_BLUE, bold=True)
        r = p.add_run(segment["text"])
        set_run_font(r, size=9.5, color=BLACK)

    out = OUT / "04_Редакционная_стенограмма.docx"
    doc.save(out)
    return out



def build_condensed_transcript() -> Path:
    source = BASE / "condensed-transcript-deepseek" / "transcript-condensed.json"
    payload = json.loads(source.read_text(encoding="utf-8"), parse_float=Decimal)
    metadata = payload["metadata"]
    segments = payload["transcript"]["segments"]
    visible_segments = [
        segment
        for segment in segments
        if segment["text"] and float(segment["start"]) >= 19 * 60
    ]

    doc = Document()
    configure_document(
        doc,
        "compact_reference_guide",
        "СОКРАЩЕННАЯ СМЫСЛОВАЯ СТЕНОГРАММА • DEEPSEEK",
    )
    add_masthead(
        doc,
        "Смысловая стенограмма",
        "Сокращенная смысловая стенограмма совещания",
        "Сокращенная производная для быстрого чтения; содержательная часть начинается с 00:19:04",
        [
            ("Модель", metadata["model"]),
            ("Видимые реплики", str(len(visible_segments))),
            ("Слова в DeepSeek-выходе", str(metadata["output_words"])),
            ("Дата формирования", "29 июля 2026 г."),
        ],
    )
    add_callout(
        doc,
        "Внимание: это сокращенная смысловая производная, а не дословная стенограмма",
        "DeepSeek V4 Flash удалил повторы, повторные формулировки и несодержательный материал. "
        "Критические факты, имена, числа, решения и принадлежность реплик необходимо проверять "
        "по документу «02_Полная_стенограмма_с_говорящими.docx» и исходной аудиозаписи.",
        fill="FFF3F3",
        color=RED,
    )
    add_table(
        doc,
        ["Показатель сокращения", "Значение"],
        [
            ["Провайдер / модель", f"{metadata['provider']} / {metadata['model']}"],
            ["Версия процедуры", metadata["prompt_version"]],
            ["Источник — артефакт 04", f"{metadata['source_nonempty_segments']} непустых сегментов / {metadata['source_words']} слов"],
            ["После объединения по говорящим", f"{metadata['merged_speaker_turns']} реплик"],
            ["Выход DeepSeek", f"{metadata['output_turns']} реплики / {metadata['output_words']} слов"],
            ["Финально показано", f"{len(visible_segments)} реплика; исключены 3 предсовещательные реплики до 00:19:00"],
            ["Сокращение против артефакта 04", f"{metadata['word_reduction_fraction'] * 100:.3f}% до финального исключения трех реплик"],
            ["Время обработки", f"{metadata['elapsed_seconds']} с"],
            ["Стоимость API", f"${metadata['api_cost_usd']}"],
        ],
        [3200, 6160],
        compact_font=9.0,
        header_fill=LIGHT_GRAY,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(
        f"Происхождение: {metadata['source_nonempty_segments']} непустых сегментов и "
        f"{metadata['source_words']} слов из артефакта 04 объединены в "
        f"{metadata['merged_speaker_turns']} реплик; DeepSeek вернул "
        f"{metadata['output_turns']} реплики и {metadata['output_words']} слов. "
        f"После исключения трех предсовещательных реплик показано {len(visible_segments)}. "
        "Текст и метки Speaker N ниже воспроизведены из transcript-condensed.json "
        "без дополнительного перефразирования."
    )
    set_run_font(r, size=9.2, color=MUTED, italic=True)

    current_bucket = -1
    for segment in visible_segments:
        start = segment["start"]
        bucket = int(start) // 600
        if bucket != current_bucket:
            start_m = bucket * 10
            end_m = start_m + 10
            p = doc.add_paragraph(
                f"{start_m // 60:02d}:{start_m % 60:02d}–{end_m // 60:02d}:{end_m % 60:02d}",
                style="Heading 2",
            )
            p.paragraph_format.page_break_before = current_bucket >= 0
            current_bucket = bucket

        p = doc.add_paragraph(style="Transcript Entry")
        p.paragraph_format.keep_together = True
        r = p.add_run(f"[{timestamp(float(start))}] ")
        set_run_font(r, size=9.2, color=MUTED)
        r = p.add_run(f"{segment['speaker']}: ")
        set_run_font(r, size=9.5, color=DARK_BLUE, bold=True)
        r = p.add_run(segment["text"])
        set_run_font(r, size=9.5, color=BLACK)

    out = OUT / "05_Сокращенная_смысловая_стенограмма.docx"
    doc.save(out)
    return out



def fmt_time(seconds: float) -> str:
    minutes, sec = divmod(seconds, 60)
    return f"{int(minutes)} мин {sec:04.1f} с"


def add_detailed_report_front_matter(doc: Document) -> None:
    doc.add_paragraph("Как работает цепочка — по шагам", style="Heading 1")
    steps = [
        "**Нормализация аудио.** Исходный MP4 приводится к единому аудиоформату, чтобы все ветви получили один и тот же сигнал.",
        "**Три независимые ветви STT.** WhisperX, MOSS и Qwen3-ASR обрабатывают одну запись независимо и дают разные чтения спорных слов, чисел и терминов.",
        "**Глобальные метки говорящих.** MOSS сначала назначает локальные метки внутри чанков. ECAPA затем объединяет 28 локальных меток в 6 глобальных анонимных кластеров. Это не установленные личности, и кластеры пока не подтверждены человеком.",
        "**Консенсус по интервалам.** 720 реплик MOSS образуют неизменяемый временной и спикерный каркас. Для каждого интервала собираются перекрывающиеся по времени кандидаты MOSS, Qwen и WhisperX.",
        "**Автоприем или арбитраж.** Явное совпадение принимается автоматически; иначе текст передается DeepSeek V4 Flash. В этом запуске: 8 автосовпадений, 712 LLM-разборов и 123 реплики с низкой уверенностью.",
        "**Ограничения DeepSeek.** Модель видит только текстовые кандидаты, не получает аудио и не может менять говорящего или временные границы. Она выбирает текст и записывает поддержку источниками, уверенность и аудит.",
        "**Каноническая стенограмма.** Выбранный текст объединяется с временными метками и глобальными кластерами MOSS. Получилось 6 175 слов; low-confidence реплики сохраняются как очередь проверки.",
        "**Два варианта протокола.** Локальная e4b и DeepSeek V4 Flash независимо извлекают темы, решения, действия, риски и вопросы с временными ссылками.",
        "**Human QA и gold-метрики.** Человек проверяет смены говорящих, числа, имена, термины и low-confidence реплики. Затем рассчитываются WER/CER, cpWER, DER/JER и метрики протокола; до этого диагностики нельзя называть фактической точностью.",
    ]
    for idx, step in enumerate(steps, 1):
        add_numbered(doc, idx, step)

    doc.add_paragraph("Материальные сильные стороны и роль каждого варианта", style="Heading 1")
    rows = [
        ["WhisperX", "Быстрый выровненный русский baseline; независимое чтение.", "209,756 с; RTF 0,04681 (≈21,4× RT).", "Нет говорящих; обязательный шлюз не пройден. Пик VRAM не записан.", "Быстрый черновик и независимый канал."],
        ["MOSS + ECAPA", "Единственный локальный маршрут здесь с глобально согласованными анонимными метками говорящих.", "28 локальных меток → 6 кластеров; пик VRAM 3,146 ГБ.", "3 119,101 с; кластеры предварительные, включая малые; нужна DER/JER-проверка.", "Выбранный структурный якорь."],
        ["Qwen3-ASR + aligner", "Независимое альтернативное чтение и временное выравнивание.", "425,287 с; RTF 0,09491; пик VRAM 10,113 ГБ.", "Нет говорящих; повторяемость 5-грамм 0,086 — риск повторов, не WER.", "Вспомогательный текстовый канал."],
        ["DeepSeek-консенсус", "Разрешает расхождения STT и оставляет поддержку, уверенность и аудит.", "6 175 слов; покрытие 99,9%; 519,006 с; $0.033694.", "Не слышит аудио; наследует diarization, время MOSS и общие ошибки STT.", "Предварительный производственный кандидат."],
        ["Локальный протокол e4b", "Приватный локальный второй взгляд; $0 API/предельной стоимости.", "59,3 с; 3 решения; 6 действий.", "Менее полный; 0% действий с явно указанным владельцем.", "Конфиденциальный fallback и второе мнение."],
        ["Протокол DeepSeek Flash", "Более полное структурированное извлечение с временными ссылками.", "47,7 с; $0.004610; 5 решений; 8 действий; покрытие зон 73,3%; владельцы 37,5%.", "Может превратить предложение в решение; текст уходит внешнему API.", "Текущий черновик протокола с обязательной human QA."],
    ]
    add_table(
        doc,
        ["Вариант", "Материально хорош в", "Факты этого запуска", "Ограничения", "Роль / статус"],
        rows,
        [1250, 1900, 2250, 2200, 1760],
        compact_font=7.6,
        header_fill=LIGHT_GRAY,
    )

    doc.add_paragraph("Почему выбрана именно эта комбинация", style="Heading 1")
    for item in [
        "**Speaker ID — жесткий шлюз.** WhisperX и Qwen сами по себе не подходят для финала: в этом запуске они не выдали глобальные метки говорящих. MOSS + ECAPA дает отсутствующую структуру.",
        "**Разнообразие ошибок.** WhisperX и Qwen оставлены как независимые чтения спорных слов, чисел и терминов; ни один канал не объявляется истинным.",
        "**DeepSeek как ограниченный текстовый арбитр.** Локальный 26B-вариант был отложен как слишком тяжелый для практического прогона на этой машине. DeepSeek V4 Flash оказался быстрым, недорогим, структурированным и поддержал checkpoint/resume; полномочия ограничены текстом.",
        "**e4b ради приватности и контроля.** Локальный протокол остается конфиденциальным fallback и независимым вторым мнением.",
        "**Выбор предварительный.** Окончательное решение возможно только после WER/CER, cpWER, DER/JER и слепого человеческого сравнения протоколов.",
    ]:
        add_bullet(doc, item)

    doc.add_paragraph("Какой режим выбирать", style="Heading 1")
    add_table(
        doc,
        ["Режим", "Цепочка", "Когда применять"],
        [
            ["Быстрый черновик", "WhisperX; при необходимости Qwen как второе чтение.", "Когда важнее скорость, а speaker ID временно не нужен."],
            ["Конфиденциально / без облака", "MOSS + ECAPA; локальная e4b для протокола.", "Когда текст нельзя передавать внешнему API; нужна ручная проверка."],
            ["Speaker-required reviewed production", "MOSS + ECAPA → multi-STT → DeepSeek-консенсус → DeepSeek-протокол → human QA.", "Текущий рабочий режим при разрешенной облачной обработке."],
            ["Максимальная уверенность", "Все три STT, 123 low-confidence реплики, ручная сверка аудио и blinded review.", "Для решений, чисел, имен, сроков и других пунктов высокого риска."],
        ],
        [2200, 3920, 3240],
        compact_font=8.7,
        header_fill=LIGHT_GRAY,
    )

    doc.add_paragraph("Происхождение данных и авторство", style="Heading 1")
    for item in [
        "**Стенограмма:** текстовые кандидаты MOSS, Qwen3-ASR и WhisperX; временной и спикерный каркас — MOSS с глобализацией меток ECAPA.",
        "**DeepSeek-консенсус:** выбирает текст при расхождениях и пишет аудит. DeepSeek получает текст, а не аудио, и не изменяет говорящих или интервалы.",
        "**Протоколы:** e4b и DeepSeek создали отдельные извлечения по канонической стенограмме. DeepSeek-версия — текущий черновик, а не подтвержденный перечень фактов.",
        "**Этот отчет:** объяснения и интерпретация архитектуры подготовлены Codex; числа взяты из JSON-артефактов запуска. Фактическая корректность текста и speaker ID пока не подтверждена human gold-разметкой.",
    ]:
        add_bullet(doc, item)



def build_report() -> Path:
    benchmark = json.loads(
        (BASE / "benchmark-final" / "benchmark.json").read_text(encoding="utf-8")
    )
    summary = json.loads(
        (BASE / "summary-comparison" / "comparison.json").read_text(encoding="utf-8")
    )
    consensus = json.loads(
        (BASE / "consensus-deepseek-flash-v2" / "transcript.json").read_text(
            encoding="utf-8"
        )
    )["metadata"]["consensus"]
    moss_meta = json.loads((BASE / "moss" / "transcript.json").read_text("utf-8"))[
        "metadata"
    ]
    qwen_meta = json.loads((BASE / "qwen" / "transcript.json").read_text("utf-8"))[
        "metadata"
    ]
    cluster_meta = json.loads(
        (BASE / "moss-global" / "transcript.json").read_text("utf-8")
    )["metadata"]["speaker_clustering"]
    refs = json.loads(
        (BASE / "human-reference" / "reference.json").read_text(encoding="utf-8")
    )["clips"]
    evaluation = json.loads(
        (BASE / "evaluation-suite" / "evaluation.json").read_text(encoding="utf-8")
    )

    doc = Document()
    configure_document(doc, "standard_business_brief", "ОЦЕНКА СИСТЕМЫ ТРАНСКРИБАЦИИ")
    add_masthead(
        doc,
        "Технический отчет",
        "Сравнение моделей транскрибации и суммаризации",
        "Запись совещания на русском языке, 74:41; обязательная идентификация говорящих",
        [
            ("Испытательная система", "Ryzen 9 9900X • 64 ГБ DDR5 • NVIDIA RTX 5070 Ti 16 ГБ"),
            ("Дата измерений", "29 июля 2026 г."),
            ("Объект", "Одна реальная запись совещания; русский язык с единичными английскими словами"),
        ],
    )
    add_callout(
        doc,
        "Главный вывод",
        "Практический текущий вариант — консенсус трех STT с временными метками и голосовыми кластерами MOSS, затем DeepSeek V4 Flash для разрешения расхождений и протокола. Это предварительный выбор: фактическая точность текста и говорящих еще не измерена, потому что человеческая эталонная разметка не заполнена.",
        fill=BLUE_GRAY,
        color=NAVY,
    )

    add_detailed_report_front_matter(doc)

    doc.add_paragraph("1. Что сравнивалось", style="Heading 1")
    add_bullet(
        doc,
        "**WhisperX large-v3:** зрелый быстрый русскоязычный базовый STT с выравниванием; в этом запуске без diarization и поэтому не удовлетворяет обязательному требованию по говорящим.",
    )
    add_bullet(
        doc,
        "**MOSS-Transcribe-Diarize:** транскрибация и нативные локальные метки говорящих; после чанков голосовые профили объединены ECAPA-кластеризацией.",
    )
    add_bullet(
        doc,
        "**Qwen3-ASR 1.7B + ForcedAligner 0.6B:** быстрый дополнительный текстовый канал с русским языком и временными метками; нативной идентификации говорящих нет.",
    )
    add_bullet(
        doc,
        "**Консенсус DeepSeek V4 Flash:** 720 временных сегментов MOSS служат якорем; тексты MOSS, Qwen и WhisperX сравниваются, говорящие сохраняются из MOSS/ECAPA.",
    )

    doc.add_paragraph("2. Измеренные операционные результаты STT", style="Heading 1")
    models = benchmark["models"]
    stt_rows = []
    model_order = ["whisperx", "moss", "qwen", "consensus"]
    vram = {
        "whisperx": "не записано",
        "moss": f"{moss_meta['peak_vram_gb']:.3f} ГБ",
        "qwen": f"{qwen_meta['peak_vram_gb']:.3f} ГБ",
        "consensus": "API; локально не требуется",
    }
    status = {
        "whisperx": "Без говорящих — не проходит",
        "moss": "6 кластеров — требует проверки",
        "qwen": "Без говорящих — вспомогательный",
        "consensus": "6 кластеров унаследованы",
    }
    for name in model_order:
        m = models[name]
        speed = 1 / m["real_time_factor"]
        stt_rows.append(
            [
                name,
                fmt_time(m["runtime_seconds"]),
                f"{m['real_time_factor']:.5f}\n({speed:.2f}× RT)",
                vram[name],
                str(m["speakers"]),
                f"${m['cost_usd']:.6f}" if m["cost_usd"] else "$0",
                status[name],
            ]
        )
    add_table(
        doc,
        ["Метод", "Время", "RTF", "Пик VRAM", "Говор.", "Стоимость", "Статус"],
        stt_rows,
        [1750, 1150, 1100, 1250, 800, 1150, 2160],
        compact_font=8.5,
        header_fill=LIGHT_GRAY,
        center_columns={1, 2, 3, 4, 5},
    )
    add_callout(
        doc,
        "Интерпретация времени",
        "RTF < 1 означает обработку быстрее длительности записи. Полная последовательная цепочка WhisperX + MOSS + ECAPA + Qwen + DeepSeek-консенсус заняла около 71 мин 17 с; с финальной DeepSeek-суммаризацией — около 72 мин 5 с. Независимые STT можно распараллелить при наличии отдельных GPU, но на одной RTX 5070 Ti они запускались последовательно.",
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run("Стоимость: ")
    set_run_font(r, bold=True)
    r = p.add_run(
        "локальные модели учитываются как $0 предельной/API-стоимости; требования к оборудованию показаны отдельно. Расход электроэнергии намеренно не рассчитывается."
    )
    set_run_font(r)

    doc.add_paragraph("3. Результат консенсуса", style="Heading 1")
    automatic_pct = 100 * consensus["automatic_agreements"] / consensus["total_units"]
    low_pct = 100 * consensus["low_confidence"] / consensus["total_units"]
    add_table(
        doc,
        ["Показатель", "Значение", "Как трактовать"],
        [
            ["Сегменты", str(consensus["total_units"]), "Якорные интервалы MOSS"],
            [
                "Автосовпадения",
                f"{consensus['automatic_agreements']} ({automatic_pct:.1f}%)",
                "Совпадения без вызова LLM",
            ],
            [
                "LLM-разногласия",
                str(consensus["llm_disputes"]),
                "Сегменты, переданные DeepSeek для выбора текста",
            ],
            [
                "Низкая уверенность",
                f"{consensus['low_confidence']} ({low_pct:.1f}%)",
                "Приоритет для ручной проверки",
            ],
            [
                "API-токены",
                f"{consensus['prompt_tokens']:,} вход / {consensus['completion_tokens']:,} выход".replace(
                    ",", " "
                ),
                f"Кэш: {consensus['prompt_cache_hit_tokens']} токена",
            ],
            ["Стоимость", "$0.033694", "Фактическая стоимость консенсусного запуска"],
        ],
        [2500, 2200, 4660],
        compact_font=9.0,
        header_fill=LIGHT_GRAY,
        center_columns={1},
    )
    p = doc.add_paragraph()
    p.add_run(
        "Кросс-модельное совпадение — диагностический сигнал, а не доказательство истины: несколько моделей могут одинаково ошибаться."
    )

    doc.add_paragraph("3.1. Автоматические диагностики без ground truth", style="Heading 2")
    diagnostic_rows = []
    for name in model_order:
        m = evaluation["transcripts"][name]
        diagnostic_rows.append(
            [
                name,
                "PASS" if m["standalone_required_gates_pass"] else "FAIL",
                f"{100 * m['timeline_coverage']:.1f}%",
                str(m["invalid_or_zero_duration_segments"]),
                f"{m['mean_pairwise_window_similarity']:.3f}",
                f"{m['repeated_5gram_fraction']:.3f}",
            ]
        )
    add_table(
        doc,
        ["Выход", "Обяз. шлюз", "Покрытие", "Плохое время", "Совпадение", "5-грамм повторы"],
        diagnostic_rows,
        [2000, 1300, 1300, 1500, 1500, 1760],
        compact_font=8.7,
        header_fill=LIGHT_GRAY,
        center_columns={1, 2, 3, 4, 5},
    )
    add_callout(
        doc,
        "Не точность",
        "Совпадение, повторы, покрытие и silhouette — диагностики, а не correctness. Повторяемость Qwen 0,086 указывает на риск зацикливания; она не измеряет WER.",
        fill="FFF8E8",
        color=GOLD,
    )

    doc.add_paragraph("4. Суммаризация: локальная e4b против DeepSeek V4 Flash", style="Heading 1")
    by_model = {entry["model"]: entry for entry in summary["models"]}
    sum_rows = []
    for model in ("gemma4:e4b", "deepseek-v4-flash"):
        m = by_model[model]
        cost = "$0" if m["provider"] == "ollama" else "$0.004610"
        sum_rows.append(
            [
                model,
                f"{m['elapsed_seconds']:.1f} с",
                f"{m['prompt_tokens']:,}".replace(",", " "),
                f"{m['completion_tokens']:,}".replace(",", " "),
                cost,
            ]
        )
    add_table(
        doc,
        ["Модель", "Время", "Входные токены", "Выходные токены", "Стоимость"],
        sum_rows,
        [2800, 1300, 1800, 1800, 1660],
        compact_font=9.2,
        header_fill=LIGHT_GRAY,
        center_columns={1, 2, 3, 4},
    )
    extraction_rows = []
    for model in ("gemma4:e4b", "deepseek-v4-flash"):
        m = by_model[model]
        extraction_rows.append(
            [
                model,
                str(m["topics"]),
                str(m["decisions"]),
                str(m["actions"]),
                str(m["risks"]),
                str(m["open_questions"]),
            ]
        )
    add_table(
        doc,
        ["Модель", "Темы", "Решения", "Действия", "Риски", "Вопросы"],
        extraction_rows,
        [2860, 1200, 1300, 1300, 1300, 1400],
        compact_font=9.2,
        header_fill=LIGHT_GRAY,
        center_columns={1, 2, 3, 4, 5},
    )
    summary_diagnostic_rows = []
    for key, label in (("e4b_summary", "gemma4:e4b"), ("deepseek_summary", "deepseek-v4-flash")):
        m = evaluation["summaries"][key]
        summary_diagnostic_rows.append(
            [
                label,
                f"{100 * m['valid_timestamp_fraction']:.1f}%",
                f"{100 * m['five_minute_source_bin_coverage']:.1f}%",
                f"{100 * m['actions_with_explicit_owner_fraction']:.1f}%",
                f"{100 * m['actions_with_explicit_due_fraction']:.1f}%",
            ]
        )
    add_table(
        doc,
        ["Модель", "Валид. ссылки", "Покрытие 5-мин зон", "Явные владельцы", "Явные сроки"],
        summary_diagnostic_rows,
        [2800, 1500, 1900, 1580, 1580],
        compact_font=8.8,
        header_fill=LIGHT_GRAY,
        center_columns={1, 2, 3, 4},
    )
    add_callout(
        doc,
        "Ограничение",
        "Большее число извлеченных пунктов, валидные временные ссылки и более широкое покрытие означают полезные диагностики, а не корректность. DeepSeek нашел 5 решений, 8 действий и 9 рисков против 3, 6 и 5 у e4b, но каждый пункт должен быть проверен по временной метке. Текущий предварительный лидер по полноте — DeepSeek V4 Flash.",
        fill="FFF8E8",
        color=GOLD,
    )

    doc.add_paragraph("5. Какие метрики обычно применяют", style="Heading 1")
    metric_rows = [
        [
            "Текст STT",
            "WER, CER",
            "Ошибка слов; ошибка символов",
            "Обязательны против человеческой дословной разметки. CER особенно полезен для русских окончаний и имен.",
        ],
        [
            "Говорящие",
            "DER, JER",
            "Пропуски речи + ложная речь + путаница говорящих; ошибка Jaccard",
            "Указать collar и правило для одновременной речи. Измеряет временную diarization.",
        ],
        [
            "Текст + говорящий",
            "cpWER / SA-WER",
            "WER после оптимального сопоставления говорящих; атрибутированный WER",
            "Ключевая метрика при обязательном Speaker ID: штрафует правильные слова у неправильного говорящего.",
        ],
        [
            "Термины и цифры",
            "NER-F1, числовая точность",
            "Имена/организации/термины; точность чисел и единиц",
            "Отдельный глоссарий русского проекта и английских терминов. Ошибки чисел имеют повышенный риск.",
        ],
        [
            "Временные метки",
            "MAE границ",
            "Средняя абсолютная ошибка начала/конца",
            "Проверять на вручную размеченных границах; отдельно — долю реплик в допустимом допуске.",
        ],
        [
            "Протокол",
            "Precision/Recall/F1",
            "Решения, действия, владельцы, сроки, риски",
            "Сопоставлять с человеческим перечнем; считать точность цитат и долю выдуманных пунктов.",
        ],
        [
            "Качество чтения",
            "Human rating",
            "Фактичность, полнота, ясность, полезность",
            "Шкала 1–5 минимум у двух проверяющих; согласованность оценщиков сообщать отдельно.",
        ],
        [
            "Эксплуатация",
            "RTF, VRAM, стоимость, отказоустойчивость",
            "Скорость, ресурс, API-затраты, успешный resume",
            "Считать на одинаковой машине и одинаковой записи; электричество не включать.",
        ],
    ]
    add_table(
        doc,
        ["Объект", "Метрика", "Что измеряет", "Правило применения"],
        metric_rows,
        [1500, 1700, 2400, 3760],
        compact_font=8.2,
        header_fill=LIGHT_GRAY,
    )

    doc.add_paragraph("6. Правило выбора без ложной точности", style="Heading 1")
    p = doc.add_paragraph()
    p.add_run(
        "Сначала применяется обязательный шлюз: метод должен выдавать глобально согласованные метки говорящих. WhisperX и Qwen в текущей конфигурации остаются вспомогательными каналами и не могут победить самостоятельно. После шлюза системы показываются как Pareto-компромисс: correctness, время, VRAM и цена остаются отдельными осями."
    )
    decision_items = [
        "**Шлюз 1 — говорящие:** глобальные speaker labels обязательны; без них самостоятельный результат дисквалифицируется.",
        "**Шлюз 2 — ground truth:** победитель по качеству определяется WER/CER, cpWER/tcpWER, DER/JER и метриками протокола только после человеческой разметки.",
        "**Pareto-выбор:** среди прошедших шлюзы отдельно сравниваются RTF, VRAM, API-стоимость и отказоустойчивость.",
        "**Критические срезы:** числа, единицы, имена, организации и единичные английские термины публикуются отдельными ошибками.",
    ]
    for idx, item in enumerate(decision_items, 1):
        add_numbered(doc, idx, item)
    add_callout(
        doc,
        "Правило",
        "Не сводить correctness, время, VRAM и цену в произвольный взвешенный балл. До появления human ground truth допустим только диагностический и операционный scorecard.",
        fill="FFF3F3",
        color=RED,
    )

    doc.add_paragraph("7. Состояние ground truth", style="Heading 1")
    total_ref_seconds = sum(float(c["end"]) - float(c["start"]) for c in refs)
    all_empty = all(not c.get("text") and not c.get("speaker_turns") for c in refs)
    add_table(
        doc,
        ["Параметр", "Текущее состояние"],
        [
            ["Подготовленные клипы", f"{len(refs)} × 30 с = {total_ref_seconds / 60:.0f} мин"],
            ["Принцип отбора", "Участки с низким межмодельным совпадением"],
            ["Дословный текст", "Не размечен" if all_empty else "Частично размечен"],
            ["Speaker turns", "Не размечены" if all_empty else "Частично размечены"],
            ["WER / CER / DER / cpWER", "Пока недоступны — требуют человеческой разметки"],
        ],
        [3000, 6360],
        compact_font=9.5,
        header_fill=LIGHT_GRAY,
    )
    p = doc.add_paragraph()
    p.add_run(
        "Для честного сравнения сложных участков текущие 6 минут полезны как challenge-set, но не представляют всю запись. Рекомендуется добавить случайную стратифицированную выборку спокойной речи, переходов говорящих, перекрытий, имен, чисел и англоязычных терминов; целевой объем — не менее 15–20 минут, размеченных двумя людьми с урегулированием расхождений."
    )
    add_callout(
        doc,
        "Политика speaker scoring",
        "Готовый permutation-invariant scorer считает DER, JER, speaker-attribution accuracy и cpWER без collar и с включенной одновременной речью. Все системы должны оцениваться с этой же политикой.",
    )
    p = doc.add_paragraph()
    p.add_run(
        "Для протоколов подготовлен blinded A/B review pack: human-reference/summary-review.json; ключ соответствия систем хранится отдельно. Рецензент сначала создает независимый gold-список решений и действий, затем оценивает поддержку, владельцев, сроки и evidence."
    )

    doc.add_paragraph("8. Предварительная рекомендация", style="Heading 1")
    recommendations = [
        "**Производственный кандидат:** MOSS + ECAPA для меток говорящих, WhisperX и Qwen как независимые чтения, DeepSeek V4 Flash как текстовый арбитр.",
        "**Быстрый черновик:** WhisperX или Qwen, если говорящие временно не нужны; для финального результата это не соответствует требованиям.",
        "**Контроль качества:** автоматически направлять все low-confidence сегменты, числа, имена и смены говорящих в очередь проверки.",
        "**Протокол:** использовать DeepSeek V4 Flash как текущий вариант по полноте, но запрещать неподтвержденные решения без временной ссылки.",
        "**Следующий шаг:** заполнить 12 клипов, посчитать WER/CER/DER/cpWER и только затем закрепить победителя численно.",
    ]
    for item in recommendations:
        add_bullet(doc, item)

    doc.add_paragraph("9. Риски и оговорки", style="Heading 1")
    risk_items = [
        "Шесть голосовых кластеров выбраны автоматически по наилучшему silhouette score 0,573; короткие первые кластеры могут быть шумом или дублями.",
        "MOSS показал самый низкий записанный пик VRAM (3,146 ГБ), но был значительно медленнее остальных STT; ранний полнофайловый режим также оказался ненадежным, поэтому использовались checkpointed chunks.",
        "Пик VRAM WhisperX в этом запуске не был записан; сравнивать его аппаратную потребность по памяти пока нельзя.",
        "Консенсус наследует ошибки временной сегментации и говорящих от MOSS; DeepSeek меняет текст, но не исправляет diarization.",
        "API-обработка передает содержание совещания внешнему провайдеру; нужен утвержденный режим работы с конфиденциальными записями.",
        "Результаты одной записи нельзя обобщать на все комнаты, микрофоны, шумы и составы участников.",
    ]
    for item in risk_items:
        add_bullet(doc, item)

    doc.add_paragraph("10. Источники и воспроизводимость", style="Heading 1")
    sources = [
        "MOSS-Transcribe-Diarize: https://huggingface.co/OpenMOSS-Team/MOSS-Transcribe-Diarize",
        "Qwen3-ASR: https://github.com/QwenLM/Qwen3-ASR",
        "WhisperX: https://github.com/m-bain/whisperX",
        "SpeechBrain ECAPA: https://huggingface.co/speechbrain/spkrec-ecapa-voxceleb",
        "MeetEval (cpWER и метрики встреч): https://github.com/fgnt/meeteval",
        "pyannote.metrics (DER/JER): https://pyannote.github.io/pyannote-metrics/",
    ]
    for source in sources:
        add_bullet(doc, source)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    r = p.add_run(
        "Версии моделей закреплены в JSON-метаданных запуска. Все времена, RTF, токены и затраты в таблицах взяты из файлов benchmark.json, comparison.json и consensus metadata."
    )
    set_run_font(r, size=9.5, color=MUTED, italic=True)

    out = OUT / "03_Сравнение_моделей_и_метрики_ПОДРОБНО.docx"
    doc.save(out)
    return out


def audit_docx(path: Path) -> dict:
    doc = Document(path)
    section = doc.sections[0]
    problems = []
    if round(section.page_width.inches, 2) != 8.5:
        problems.append("page width")
    if round(section.page_height.inches, 2) != 11:
        problems.append("page height")
    for attr in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        if round(getattr(section, attr).inches, 2) != 1.0:
            problems.append(attr)
    for table_idx, table in enumerate(doc.tables):
        grid_widths = [
            int(col.get(qn("w:w"))) for col in table._tbl.tblGrid.gridCol_lst
        ]
        if sum(grid_widths) != CONTENT_WIDTH_DXA:
            problems.append(f"table {table_idx} grid total {sum(grid_widths)}")
        for row_idx, row in enumerate(table.rows):
            widths = [
                int(cell._tc.get_or_add_tcPr().find(qn("w:tcW")).get(qn("w:w")))
                for cell in row.cells
            ]
            if widths != grid_widths:
                problems.append(f"table {table_idx} row {row_idx} tcW mismatch")
    return {
        "path": str(path),
        "paragraphs": len(doc.paragraphs),
        "tables": len(doc.tables),
        "problems": problems,
    }


def main() -> None:
    paths = [build_minutes(), build_transcript(), build_report()]
    report = [audit_docx(path) for path in paths]
    qa_path = OUT / "structural-audit.json"
    qa_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    for item in report:
        print(json.dumps(item, ensure_ascii=True))


if __name__ == "__main__":
    main()
