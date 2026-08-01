from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt

from build_minutes_evaluation_docx import (
    BLUE,
    GREEN,
    INK,
    MUTED,
    PALE,
    WARN,
    add_callout,
    configure,
    set_font,
    set_table_geometry,
    shade,
)


LABELS = {
    "deepseek-v4-flash": "DeepSeek V4 Flash",
    "gemma4:e4b": "Gemma 4 e4b",
    "gemma4:12b": "Gemma 4 12B",
    "gemma4:26b": "Gemma 4 26B",
    "gemma4:31b": "Gemma 4 31B",
}

RUN_METRICS = {
    "deepseek-v4-flash": Path(
        "outputs/production-protocol-universal-v6-integrated-rep2/"
        "deepseek-v4-flash/metrics.json"
    ),
    "gemma4:e4b": Path(
        "outputs/production-protocol-universal-v6-local/gemma4-e4b/metrics.json"
    ),
    "gemma4:12b": Path(
        "outputs/production-protocol-universal-v6-local-rolefix/gemma4-12b/metrics.json"
    ),
    "gemma4:26b": Path(
        "outputs/production-protocol-universal-v6-26b/gemma4-26b/metrics.json"
    ),
    "gemma4:31b": Path(
        "outputs/production-protocol-universal-v6-31b/gemma4-31b/metrics.json"
    ),
}

GENERALIZATION = {
    "deepseek-v4-flash": Path(
        "outputs/protocol-generalization-v6/deepseek-v4-flash/"
        "universal-v6-action-recall/result.json"
    ),
    "gemma4:e4b": Path(
        "outputs/protocol-generalization-v6/gemma4-e4b/"
        "universal-v6-final/result.json"
    ),
    "gemma4:12b": Path(
        "outputs/protocol-generalization-v6-12b-rolefix/gemma4-12b/"
        "universal-v6-12b-rolefix/result.json"
    ),
    "gemma4:26b": Path(
        "outputs/protocol-generalization-v6/gemma4-26b/"
        "universal-v6-final-16k/result.json"
    ),
    "gemma4:31b": Path(
        "outputs/protocol-generalization-v6/gemma4-31b/"
        "universal-v6-final-16k/result.json"
    ),
}

HARDWARE = {
    "deepseek-v4-flash": "API; локальная VRAM не требуется",
    "gemma4:e4b": "9,6 ГБ файла модели; полностью помещается в 16 ГБ VRAM",
    "gemma4:12b": "7,6 ГБ файла модели; полностью помещается в 16 ГБ VRAM",
    "gemma4:26b": "18 ГБ resident; 73% GPU / 27% CPU; VRAM во время работы ≈15,3 ГБ",
    "gemma4:31b": "21 ГБ resident; 54% GPU / 46% CPU; VRAM во время работы ≈15,3 ГБ",
}

MANUAL = {
    "deepseek-v4-flash": (
        "Самые осторожные и редакторски естественные формулировки. Все три итоговых "
        "поручения переданы верно; лучше других различает обсуждение, статус и решение.",
        "Сжимает сильнее локальных моделей: в выбранной финальной версии 8 пунктов "
        "обсуждения против 13 у человеческого документа. Требует отправки текста в API.",
        "Готовый протокол секретаря: короткий, уверенный, без ощущения машинной стенограммы.",
    ),
    "gemma4:e4b": (
        "Быстро работает полностью локально и удерживает базовую форму документа.",
        "Потеряла большую часть содержания, исказила названия и смысл отдельных поручений; "
        "два пункта обсуждения недостаточны для рабочего протокола.",
        "Черновой конспект, который нельзя выдавать коллегам без переписывания.",
    ),
    "gemma4:12b": (
        "Лучшее локальное покрытие: 11 содержательных пунктов и три ключевых поручения. "
        "После персональной подсказки корректно принимает роли и подразделения как владельцев.",
        "Иногда превращает обсуждаемую возможность в установленный факт: например, "
        "«смена локации исключена». Есть риск слишком уверенно склеить бюджетные сценарии.",
        "Старательный технический секретарь: полезный и подробный, но нуждается в финальной вычитке.",
    ),
    "gemma4:26b": (
        "Хорошая структура, аккуратная русская терминология и меньше конфликтующих чисел, "
        "чем у e4b. После верификатора оставляет три правильных поручения.",
        "Формулировки часто слишком общие и теряют полезные числа. До фильтра добавила "
        "дублирующую встречу с руководством. Медленнее 12B без устойчивого выигрыша.",
        "Солидный, но осторожный референт: звучит хорошо, сообщает меньше, чем ожидаешь от размера.",
    ),
    "gemma4:31b": (
        "Чистый технический русский, аккуратные числовые факты и три проверяемых "
        "поручения после детерминированного фильтра.",
        "Потеряла явное поручение поговорить с Министерством и оставила только шесть "
        "пунктов обсуждения. До фильтра назначила упомянутого эксперта ответственным и "
        "превратила запланированную встречу с Росатомом в новое поручение.",
        "Полированный, но слишком дорогой по времени технический меморандум; заметного "
        "практического превосходства над 12B или 26B нет.",
    ),
}


def load_json(path: Path) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


def fmt(value: float) -> str:
    return f"{value:.3f}".replace(".", ",")


def seconds(value: float) -> str:
    if value < 180:
        return f"{value:.1f} с".replace(".", ",")
    minutes = int(value // 60)
    remainder = int(round(value - 60 * minutes))
    return f"{minutes} мин {remainder:02d} с"


def repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def add_title(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(4)
    set_font(
        p.add_run("Сравнение моделей для протокола совещания"),
        size=23,
        bold=True,
        color=INK,
    )
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    set_font(
        p.add_run(
            "DeepSeek V4 Flash и локальные Gemma 4: метрики, ручная вычитка, "
            "чувствительность к промпту, время и аппаратные требования"
        ),
        size=13,
        color=MUTED,
    )
    for label, value in (
        ("Дата измерений", "31.07.2026"),
        ("Система", "Ryzen 9 9900X · 64 ГБ DDR5 · RTX 5070 Ti 16 ГБ"),
        ("Материал", "русскоязычное рабочее совещание, 1 ч 14 мин 39 с, отдельные английские термины"),
        ("Режим", "температура 0; распознавание и идентификация говорящих выполнены отдельно"),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        set_font(p.add_run(f"{label}: "), size=10.5, bold=True)
        set_font(p.add_run(value), size=10.5)


def add_executive_decision(doc: Document) -> None:
    add_callout(
        doc,
        "Рекомендация",
        "DeepSeek V4 Flash — основной режим для несекретных встреч: лучший редакторский "
        "контроль и минимальное время. Gemma 4 12B — основной полностью локальный режим: "
        "максимальное покрытие при приемлемой скорости, но с обязательным "
        "детерминированным фильтром и короткой вычиткой секретаря. 26B оставлять в "
        "Advanced; 31B оценивать только как аварийный офлайн-вариант с CPU-offload. "
        "e4b не использовать для финального протокола.",
        GREEN,
    )


def add_metric_notes(doc: Document) -> None:
    doc.add_heading("1. Что именно измеряется", level=1)
    doc.add_paragraph(
        "Human similarity — близость к документу, который подготовил человек. Это "
        "полезный стилевой ориентир, но не абсолютная истина: человеческий протокол "
        "содержит ФИО, точные сроки и сведения из рабочих материалов, отсутствующие в аудио."
    )
    doc.add_paragraph(
        "Transcript-grounded — покрытие только тех пунктов человеческого документа, "
        "которые можно восстановить по стенограмме. Это не полноценная проверка истинности "
        "каждого предложения: поэтому численные оценки дополнены ручной вычиткой."
    )
    doc.add_paragraph(
        "Discuss/Resolve F1 — покрытие обсуждений и поручений. Numbers — совпадение чисел "
        "только в автоматически сопоставленных пунктах; ноль может означать отсутствие "
        "сопоставленной числовой пары, а не обязательно числовую ошибку."
    )


def add_current_quality(
    doc: Document,
    verified: dict[str, Any],
    run_metrics: dict[str, dict[str, Any]],
    generalization: dict[str, dict[str, Any]],
) -> None:
    doc.add_heading("2. Финальные версии после универсального v6-фильтра", level=1)
    doc.add_paragraph(
        "В таблице оценивается именно документ, который попадёт в Word после "
        "детерминированной проверки поручений. Все модели получили персональную короткую "
        "надстройку к одному универсальному промпту."
    )
    table = doc.add_table(rows=1, cols=8)
    table.style = "Table Grid"
    headers = ["Модель", "Human", "Grounded", "Discuss F1", "Resolve F1", "Числа", "D/R", "Universal"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=7.6, bold=True)
    repeat_header(table.rows[0])
    by_model = {item["model"]: item for item in verified["runs"]}
    for model in LABELS:
        row = by_model.get(model)
        if not row:
            continue
        gen = generalization.get(model, {})
        values = [
            LABELS[model],
            fmt(row["human_similarity_score"]),
            fmt(row["transcript_grounded_score"]),
            fmt(row["discussed_vs_human"]["f1"]),
            fmt(row["resolved_vs_human"]["f1"]),
            fmt(row["numeric_fidelity_on_matched_items"]),
            f"{row['discussed_count']}/{row['resolved_count']}",
            f"{gen.get('passed', 0)}/{gen.get('total', 0)}",
        ]
        for cell, value in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(value), size=7.4)
    set_table_geometry(table, [1650, 900, 1030, 1050, 1050, 800, 780, 2100])
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    set_font(
        p.add_run(
            "Universal = 60 проверок на 8 коротких контрольных случаях: поручение vs "
            "упоминание, conflicting numbers, RU/EN, brainstorm, реальное поручение по "
            "протоколу и переименованные сущности."
        ),
        size=9,
        color=MUTED,
    )


def add_operation_table(
    doc: Document,
    run_metrics: dict[str, dict[str, Any]],
) -> None:
    doc.add_heading("3. Время, стоимость и требования к железу", level=1)
    deepseek_time = run_metrics["deepseek-v4-flash"]["elapsed_seconds"]
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    headers = ["Модель", "Полный цикл LLM", "Отн. к DeepSeek", "Стоимость", "Размещение модели"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=8, bold=True)
    repeat_header(table.rows[0])
    for model in LABELS:
        metrics = run_metrics.get(model)
        if not metrics:
            continue
        if model == "deepseek-v4-flash":
            hit = metrics.get("prompt_cache_hit_tokens", 0)
            miss = metrics.get("prompt_cache_miss_tokens", metrics.get("prompt_tokens", 0))
            output = metrics.get("completion_tokens", 0)
            cost = hit / 1e6 * 0.0028 + miss / 1e6 * 0.14 + output / 1e6 * 0.28
            cost_text = f"${cost:.4f}"
        else:
            cost_text = "локально: $0 API"
        values = [
            LABELS[model],
            seconds(metrics["elapsed_seconds"]),
            f"{metrics['elapsed_seconds'] / deepseek_time:.1f}×".replace(".", ","),
            cost_text,
            HARDWARE[model],
        ]
        for cell, value in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(value), size=7.8)
    set_table_geometry(table, [1500, 1350, 1150, 1250, 4110])
    doc.add_paragraph(
        "Электроэнергия намеренно не оценивается. Для локальных моделей учитываются "
        "время, размер модели и необходимость CPU/RAM offload. Цена DeepSeek рассчитана "
        "по фактическим cache hit/miss/output tokens этого запуска."
    )


def add_manual_read(doc: Document) -> None:
    doc.add_heading("4. Ручная вычитка и «вайб» результата", level=1)
    doc.add_paragraph(
        "Автоматический рейтинг не определяет победителя. Ниже — содержательная оценка "
        "финального синтеза каждой модели после чтения целиком."
    )
    for model in LABELS:
        strengths, weaknesses, vibe = MANUAL[model]
        doc.add_heading(LABELS[model], level=2)
        for label, text in (
            ("Сильные стороны", strengths),
            ("Риски", weaknesses),
            ("Впечатление", vibe),
        ):
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(4)
            set_font(p.add_run(label + ". "), size=10.7, bold=True, color=BLUE)
            set_font(p.add_run(text), size=10.7)


def add_sensitivity(doc: Document, matrix: dict[str, Any]) -> None:
    doc.add_heading("5. Как шлифовались промпты и где наступило плато", level=1)
    doc.add_paragraph(
        "Итерации не сводились к одному промпту для всех. Сначала использовались общие "
        "P0–P3, затем ответ каждой модели читался целиком, фиксировалась конкретная ошибка "
        "и добавлялась минимальная персональная поправка. Неудачные версии не удалялись."
    )
    doc.add_paragraph(
        "DeepSeek получил проверку статуса против решения и редакторское сжатие; e4b — "
        "короткие русские фразы и запрет догадок; 12B — признание ролей/подразделений "
        "валидными владельцами; 26B — раздельные числовые факты и запрет вычислять новые "
        "диапазоны; 31B — агрессивное сжатие и строгую привязку срока к поручению."
    )
    by_model: dict[str, list[dict[str, Any]]] = {}
    for item in matrix["runs"]:
        by_model.setdefault(item["model"], []).append(item)
    table = doc.add_table(rows=1, cols=7)
    table.style = "Table Grid"
    headers = ["Модель", "Итераций", "Human min–max", "Лучший Human", "Grounded min–max", "Лучший Grounded", "Вывод"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=7.5, bold=True)
    repeat_header(table.rows[0])
    conclusions = {
        "deepseek-v4-flash": "P3 лучше; дальнейшее усложнение колеблется",
        "gemma4:e4b": "сильно чувствительна; tuned легко ухудшает",
        "gemma4:12b": "P10 повышает покрытие; нужен field guard",
        "gemma4:26b": "разные профили меняют абстрактность, не дают лидерства",
        "gemma4:31b": "P7 лучший in-sample, но цена по времени высока",
    }
    for model in LABELS:
        rows = by_model.get(model, [])
        if not rows:
            continue
        best_h = max(rows, key=lambda item: item["human_similarity_score"])
        best_g = max(rows, key=lambda item: item["transcript_grounded_score"])
        values = [
            LABELS[model],
            str(len(rows)),
            f"{fmt(min(x['human_similarity_score'] for x in rows))}–{fmt(max(x['human_similarity_score'] for x in rows))}",
            f"{best_h['profile']} ({fmt(best_h['human_similarity_score'])})",
            f"{fmt(min(x['transcript_grounded_score'] for x in rows))}–{fmt(max(x['transcript_grounded_score'] for x in rows))}",
            f"{best_g['profile']} ({fmt(best_g['transcript_grounded_score'])})",
            conclusions[model],
        ]
        for cell, value in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(value), size=7.1)
    set_table_geometry(table, [1300, 650, 1250, 1700, 1350, 1750, 1360])
    add_callout(
        doc,
        "Критерий плато",
        "Универсальный v6 принят только после 60/60 на контрольных случаях у каждой "
        "модели. После этого изменения, улучшающие одну запись без нового независимого "
        "материала, считаются риском overfit. На рабочем совещании остаётся ручная "
        "диагностика, но она больше не используется как единственный оптимизируемый score.",
        WARN,
    )


def add_full_sensitivity_appendix(doc: Document, matrix: dict[str, Any]) -> None:
    doc.add_page_break()
    doc.add_heading("Приложение A. Полная матрица чувствительности", level=1)
    doc.add_paragraph(
        "Все сохранённые итерации на исходной встрече. Более высокий score не гарантирует "
        "лучший документ: ручная вычитка и универсальный контроль обязательны."
    )
    table = doc.add_table(rows=1, cols=6)
    table.style = "Table Grid"
    headers = ["Модель", "Профиль", "Human", "Grounded", "Числа", "Время"]
    for cell, text in zip(table.rows[0].cells, headers):
        shade(cell, PALE)
        set_font(cell.paragraphs[0].add_run(text), size=8, bold=True)
    repeat_header(table.rows[0])
    for item in sorted(matrix["runs"], key=lambda row: (row["model"], row["profile"])):
        values = [
            LABELS.get(item["model"], item["model"]),
            item["profile"],
            fmt(item["human_similarity_score"]),
            fmt(item["transcript_grounded_score"]),
            fmt(item["numeric_fidelity_on_matched_items"]),
            seconds(item["elapsed_seconds"]),
        ]
        for cell, value in zip(table.add_row().cells, values):
            set_font(cell.paragraphs[0].add_run(value), size=7.4)
    set_table_geometry(table, [1650, 2600, 1000, 1100, 950, 2060])


def add_architecture_and_recommendations(doc: Document) -> None:
    doc.add_heading("6. Что детерминировано, а что остаётся стохастичным", level=1)
    doc.add_paragraph(
        "Стохастичны STT, идентификация говорящих и LLM-синтез. Температура 0 уменьшает "
        "вариативность, но не превращает генерацию в математически детерминированную."
    )
    doc.add_paragraph(
        "Детерминированны разбиение на фрагменты, формат временных меток, JSON-схема, "
        "проверка evidence, запрет неподтверждённых владельцев/сроков/чисел, фильтрация "
        "служебной болтовни, дедупликация поручений и генерация настоящей таблицы Word."
    )
    doc.add_heading("7. Практическая конфигурация приложения", level=1)
    for title, text in (
        (
            "По умолчанию",
            "DeepSeek V4 Flash для разрешённых к отправке текстов; Gemma 4 12B для "
            "конфиденциального локального режима.",
        ),
        (
            "Расширенный режим",
            "26B можно показывать как экспериментальный вариант. 31B — только с ясным "
            "предупреждением о CPU-offload и времени. e4b скрыть из списка финального протокола.",
        ),
        (
            "Карточка встречи",
            "Дата, место, ФИО и утверждённая повестка вводятся отдельно секретарём. "
            "Модель не должна угадывать эти поля из речи.",
        ),
        (
            "Word",
            "DOCX создаётся автоматически только из проверенного protocol.json. Таблица "
            "поручений имеет фиксированную геометрию, повторяемый заголовок и адаптивную высоту строк.",
        ),
    ):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        set_font(p.add_run(title + ". "), size=11, bold=True, color=BLUE)
        set_font(p.add_run(text), size=11)


def add_limitations(doc: Document) -> None:
    doc.add_heading("8. Ограничения", level=1)
    doc.add_paragraph(
        "Это глубокий case study одной встречи, а не окончательный рейтинг для всех "
        "совещаний. Универсальные синтетические случаи уменьшают риск overfit, но не "
        "заменяют набор независимых реальных записей."
    )
    doc.add_paragraph(
        "Speaker ID обязателен для полной стенограммы, но анонимные Speaker N намеренно "
        "не переносятся в владельцев поручений. Оценка diarization должна вестись отдельно "
        "по DER/JER на размеченных фрагментах."
    )
    doc.add_paragraph(
        "Тариф DeepSeek может измениться. В расчёте использованы официальные цены на "
        "31.07.2026: cache hit $0,0028, cache miss $0,14 и output $0,28 за 1 млн токенов."
    )


def build(
    verified_path: Path,
    matrix_path: Path,
    output_path: Path,
) -> None:
    project = Path(__file__).resolve().parents[1]
    verified = load_json(verified_path)
    matrix = load_json(matrix_path)
    run_metrics = {
        model: load_json(project / path)
        for model, path in RUN_METRICS.items()
        if (project / path).is_file()
    }
    generalization = {
        model: load_json(project / path)
        for model, path in GENERALIZATION.items()
        if (project / path).is_file()
    }
    doc = Document()
    configure(doc)
    footer = doc.sections[0].footer.paragraphs[0]
    for run in list(footer.runs):
        footer._p.remove(run._r)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(
        footer.add_run("Внутренний рабочий отчёт · 31.07.2026"),
        size=8.5,
        color=MUTED,
    )
    add_title(doc)
    add_executive_decision(doc)
    add_metric_notes(doc)
    add_current_quality(doc, verified, run_metrics, generalization)
    add_operation_table(doc, run_metrics)
    add_manual_read(doc)
    add_sensitivity(doc, matrix)
    add_architecture_and_recommendations(doc)
    add_limitations(doc)
    add_full_sensitivity_appendix(doc, matrix)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--verified", type=Path, required=True)
    parser.add_argument("--matrix", type=Path, required=True)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    build(args.verified.resolve(), args.matrix.resolve(), args.output.resolve())
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
