from __future__ import annotations

from docx import Document

from build_minutes_evaluation_docx import BLUE, set_font


def add_docx_automation(doc: Document) -> None:
    doc.add_heading("Автоматическая генерация Word в приложении", level=1)
    steps = [
        (
            "Обязательный этап.",
            "После распознавания, определения говорящих, консенсуса, редакционной и "
            "сокращённой версий pipeline сам запускает BUILDING_DOCX. Кнопка или ручной "
            "экспорт не требуются.",
        ),
        (
            "Атомарное сохранение.",
            "Каждый документ сначала записывается во временный .docx, проверяется как "
            "целый ZIP/OOXML с непустым видимым текстом и только потом заменяет итоговый файл.",
        ),
        (
            "Статус COMPLETE только после 4/4.",
            "Pipeline проверяет точные имена, расположение, целостность и текст всех четырёх "
            "файлов. Повреждённый, пустой, перемещённый или отсутствующий DOCX заставляет "
            "этап пересобраться, а не выдаёт ложное «Готово».",
        ),
        (
            "Проверено.",
            "59 автоматических тестов, отдельный document smoke-тест и реальное задание "
            "в кириллическом OneDrive-пути: четыре из четырёх документов успешно валидированы.",
        ),
    ]
    for label, text in steps:
        p = doc.add_paragraph()
        set_font(p.add_run(label + " "), size=10.8, bold=True, color=BLUE)
        set_font(p.add_run(text), size=10.8)
