from __future__ import annotations

import argparse
import json
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .document_validation import valid_docx
from .jobs import DOCUMENTS


ACCENT = RGBColor(43, 94, 87)
INK = RGBColor(31, 41, 51)
MUTED = RGBColor(94, 105, 115)
PALE = "EAF3F1"


def _timestamp(seconds: float) -> str:
    total = max(0, int(float(seconds)))
    hours, remainder = divmod(total, 3600)
    minutes, secs = divmod(remainder, 60)
    return f"{hours:02d}:{minutes:02d}:{secs:02d}"


def _shade(paragraph, fill: str = PALE) -> None:
    properties = paragraph._p.get_or_add_pPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _configure(document: Document, *, title: str, subtitle: str) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.0)

    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 26), ("Heading 1", 17), ("Heading 2", 13)):
        style = styles[name]
        style.font.name = "Aptos Display"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
        style.font.size = Pt(size)
        style.font.color.rgb = ACCENT
    title_p = document.add_paragraph(style="Title")
    title_p.add_run(title)
    subtitle_p = document.add_paragraph()
    subtitle_p.paragraph_format.space_after = Pt(18)
    run = subtitle_p.add_run(subtitle)
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED

    header = section.header.paragraphs[0]
    header.text = "Meeting Documents"
    header.runs[0].font.name = "Aptos"
    header.runs[0].font.size = Pt(8.5)
    header.runs[0].font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Создано локальным приложением Meeting Documents  •  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED


def _add_notice(document: Document, text: str, *, warning: bool = False) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.35)
    paragraph.paragraph_format.right_indent = Cm(0.35)
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(14)
    _shade(paragraph, "FFF4E5" if warning else PALE)
    run = paragraph.add_run(text)
    run.bold = True
    run.font.color.rgb = RGBColor(128, 78, 15) if warning else ACCENT


def _load_transcript(path: Path) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    transcript = payload.get("transcript", payload)
    segments = transcript.get("segments", [])
    if not isinstance(segments, list) or not segments:
        raise ValueError(f"В {path.name} нет текста стенограммы.")
    return payload, [item for item in segments if str(item.get("text", "")).strip()]


def _source_subtitle(job: dict[str, Any], *, kind: str) -> str:
    filename = job["input"]["filename"]
    model = job["text_processing"]["model"]
    return f"{filename}  •  {kind}  •  текстовая модель: {model}"


def _add_transcript_entries(
    document: Document, segments: Iterable[dict[str, Any]]
) -> None:
    for segment in segments:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.space_after = Pt(4)
        start = _timestamp(float(segment.get("start", 0)))
        end = _timestamp(float(segment.get("end", segment.get("start", 0))))
        speaker = str(segment.get("speaker") or "Speaker")
        label = paragraph.add_run(f"[{start}–{end}]  {speaker}: ")
        label.bold = True
        label.font.color.rgb = ACCENT
        paragraph.add_run(re.sub(r"\s+", " ", str(segment["text"])).strip())


def _table_cells(line: str) -> list[str]:
    value = line.strip().strip("|")
    return [
        _strip_markdown(cell.replace("\\|", "|").strip())
        for cell in re.split(r"(?<!\\)\|", value)
    ]


def _is_table_separator(line: str) -> bool:
    cells = _table_cells(line)
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _set_cell_margins(cell, *, top: int, start: int, bottom: int, end: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    margins = properties.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        properties.append(margins)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    properties = table._tbl.tblPr
    table_width = properties.first_child_found_in("w:tblW")
    if table_width is None:
        table_width = OxmlElement("w:tblW")
        properties.insert(0, table_width)
    table_width.set(qn("w:w"), str(sum(widths)))
    table_width.set(qn("w:type"), "dxa")
    indent = properties.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        properties.append(indent)
    indent.set(qn("w:w"), "120")
    indent.set(qn("w:type"), "dxa")
    layout = properties.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        properties.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        column = OxmlElement("w:gridCol")
        column.set(qn("w:w"), str(width))
        grid.append(column)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_width = cell._tc.get_or_add_tcPr().get_or_add_tcW()
            cell_width.set(qn("w:w"), str(width))
            cell_width.set(qn("w:type"), "dxa")
            _set_cell_margins(cell, top=80, start=120, bottom=80, end=120)


def _add_markdown_table(
    document: Document,
    headers: list[str],
    rows: list[list[str]],
) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    widths = [5760, 1800, 1800] if len(headers) == 3 else [9360 // len(headers)] * len(headers)
    widths[-1] += 9360 - sum(widths)
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        cell.text = value
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), PALE)
        cell._tc.get_or_add_tcPr().append(shading)
        for run in cell.paragraphs[0].runs:
            run.bold = True
            run.font.color.rgb = ACCENT
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    for values in rows:
        cells = table.add_row().cells
        padded = [*values, *([""] * len(headers))][: len(headers)]
        for index, value in enumerate(padded):
            cells[index].text = value
            cells[index].paragraphs[0].alignment = (
                WD_ALIGN_PARAGRAPH.LEFT
                if index == 0
                else WD_ALIGN_PARAGRAPH.CENTER
            )
    _set_table_geometry(table, widths)
    for row in table.rows:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_before = Pt(0)
                paragraph.paragraph_format.space_after = Pt(3)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Aptos"
                    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
                    run.font.size = Pt(9.5)
    document.add_paragraph().paragraph_format.space_after = Pt(2)


def _add_markdown(document: Document, text: str) -> None:
    lines = text.splitlines()
    index = 0
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if (
            index + 1 < len(lines)
            and "|" in line
            and _is_table_separator(lines[index + 1].strip())
        ):
            headers = _table_cells(line)
            rows: list[list[str]] = []
            index += 2
            while index < len(lines) and "|" in lines[index]:
                rows.append(_table_cells(lines[index]))
                index += 1
            _add_markdown_table(document, headers, rows)
            continue
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            level = min(2, len(heading.group(1)))
            document.add_paragraph(heading.group(2), style=f"Heading {level}")
            index += 1
            continue
        bullet = re.match(r"^[-*]\s+(.+)$", line)
        if bullet:
            document.add_paragraph(_strip_markdown(bullet.group(1)), style="List Bullet")
            index += 1
            continue
        numbered = re.match(r"^\d+[.)]\s+(.+)$", line)
        if numbered:
            document.add_paragraph(
                _strip_markdown(numbered.group(1)), style="List Number"
            )
            index += 1
            continue
        quote = re.match(r"^>\s*(.+)$", line)
        if quote:
            _add_notice(document, _strip_markdown(quote.group(1)), warning=True)
            index += 1
            continue
        document.add_paragraph(_strip_markdown(line))
        index += 1


def _strip_markdown(text: str) -> str:
    text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
    text = re.sub(r"`(.+?)`", r"\1", text)
    return text


def build_minutes(job_dir: Path, job: dict[str, Any]) -> Path:
    model = job["text_processing"]["model"]
    slug = re.sub(r"[^0-9A-Za-zА-Яа-я._-]+", "-", model).strip("-").lower()
    source = job_dir / "stages" / "minutes" / f"minutes-{slug}.md"
    if not source.is_file():
        candidates = list((job_dir / "stages" / "minutes").glob("minutes-*.md"))
        if not candidates:
            raise FileNotFoundError("Не найден подготовленный протокол.")
        source = candidates[0]
    document = Document()
    _configure(
        document,
        title="Протокол совещания",
        subtitle=_source_subtitle(job, kind="структурированная версия"),
    )
    _add_notice(
        document,
        "Проверьте критические решения, сроки, имена и числа по полной стенограмме и аудиозаписи.",
        warning=True,
    )
    _add_markdown(document, source.read_text(encoding="utf-8"))
    return _save(document, job_dir, DOCUMENTS[0]["filename"])


def build_transcript(
    job_dir: Path,
    job: dict[str, Any],
    *,
    source: Path,
    title: str,
    filename: str,
    kind: str,
    notice: str,
    warning: bool,
) -> Path:
    _, segments = _load_transcript(source)
    document = Document()
    _configure(document, title=title, subtitle=_source_subtitle(job, kind=kind))
    _add_notice(document, notice, warning=warning)
    _add_transcript_entries(document, segments)
    return _save(document, job_dir, filename)


def _save(document: Document, job_dir: Path, filename: str) -> Path:
    output = job_dir / "documents" / filename
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_name(f".{output.stem}.{os.getpid()}.tmp.docx")
    try:
        document.save(temporary)
        if not valid_docx(temporary, expected_parent=output.parent):
            raise ValueError(f"Создан повреждённый или пустой Word-файл: {filename}")
        temporary.replace(output)
    finally:
        temporary.unlink(missing_ok=True)
    return output


def build_all(job_dir: Path) -> list[dict[str, Any]]:
    job = json.loads((job_dir / "job.json").read_text(encoding="utf-8"))
    outputs = [
        build_minutes(job_dir, job),
        build_transcript(
            job_dir,
            job,
            source=job_dir / "stages" / "consensus" / "transcript.json",
            title="Полная стенограмма",
            filename=DOCUMENTS[1]["filename"],
            kind="полная версия",
            notice=(
                "Сохранены временные метки и анонимные метки говорящих. "
                "Текст сверен по нескольким системам распознавания."
            ),
            warning=False,
        ),
        build_transcript(
            job_dir,
            job,
            source=job_dir / "stages" / "cleaned" / "transcript-cleaned.json",
            title="Редакционная стенограмма",
            filename=DOCUMENTS[2]["filename"],
            kind="редакционная версия",
            notice=(
                "Не дословная версия: удалены речевой мусор, проверки связи, "
                "ложные старты и непосредственные повторы."
            ),
            warning=True,
        ),
        build_transcript(
            job_dir,
            job,
            source=job_dir / "stages" / "condensed" / "transcript-condensed.json",
            title="Сокращённая смысловая стенограмма",
            filename=DOCUMENTS[3]["filename"],
            kind="сокращённая версия",
            notice=(
                "Не дословная версия: оставлено только существенное. "
                "Порядок реплик, говорящие и временные метки сохранены."
            ),
            warning=True,
        ),
    ]
    manifest = []
    for definition, path in zip(DOCUMENTS, outputs):
        manifest.append(
            {
                **definition,
                "path": str(path.resolve()),
                "size_bytes": path.stat().st_size,
            }
        )
    (job_dir / "documents" / "manifest.json").write_text(
        json.dumps(manifest, ensure_ascii=False, indent=2) + "\n",
        encoding="utf-8",
    )
    return manifest


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Build the four clerk-facing DOCX files.")
    parser.add_argument("--job-dir", type=Path, required=True)
    args = parser.parse_args(argv)
    manifest = build_all(args.job_dir.expanduser().resolve())
    print(json.dumps({"documents": manifest}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
